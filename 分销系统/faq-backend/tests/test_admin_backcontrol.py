"""Tests for back-office account, ranking, and participation report APIs."""

from __future__ import annotations

from datetime import datetime
from io import BytesIO

from openpyxl import load_workbook
import pytest
from sqlalchemy import select

from app.core.config import settings
from app.models.energy import EnergyRedemptionRecord, EnergyTransaction
from app.models.lottery import LotteryDraw, LotteryWinner
from app.models.monthly import MonthlyRankSnapshot
from app.models.record import AnswerRecord
from app.models.user import User


def admin_headers(token: str) -> dict[str, str]:
    return {"Authorization": f"Bearer {token}"}


@pytest.fixture
def admin_credentials():
    old_username = getattr(settings, "ADMIN_USERNAME", None)
    old_password = getattr(settings, "ADMIN_PASSWORD", None)
    old_secret = getattr(settings, "ADMIN_SECRET_KEY", None)

    settings.ADMIN_USERNAME = "admin"
    settings.ADMIN_PASSWORD = "secret123"
    settings.ADMIN_SECRET_KEY = "admin-secret"

    yield {"username": "admin", "password": "secret123"}

    settings.ADMIN_USERNAME = old_username
    settings.ADMIN_PASSWORD = old_password
    settings.ADMIN_SECRET_KEY = old_secret


async def login_admin(client, credentials) -> str:
    response = await client.post("/api/admin/auth/login", json=credentials)
    assert response.status_code == 200
    return response.json()["data"]["token"]


@pytest.mark.asyncio
async def test_admin_can_manage_user_accounts(client, admin_credentials):
    token = await login_admin(client, admin_credentials)
    headers = admin_headers(token)

    create_resp = await client.post(
        "/api/admin/users",
        headers=headers,
        json={
            "phone": "13800001111",
            "login_username": "sales_001",
            "login_password": "pass123",
            "real_name": "张三",
            "nickname": "",
            "province": "上海",
            "company": "测试公司",
            "job_role": "销售",
            "profile_verified": True,
        },
    )

    assert create_resp.status_code == 200
    created = create_resp.json()["data"]
    assert created["real_name"] == "张三"
    assert created["job_role"] == "销售"
    assert created["job_role_label"] == "销售"
    assert created["company"] == "测试公司"

    list_resp = await client.get("/api/admin/users", headers=headers, params={"company": "测试公司"})
    assert list_resp.status_code == 200
    rows = list_resp.json()["data"]["items"]
    assert [row["id"] for row in rows] == [created["id"]]
    assert list_resp.json()["data"]["total"] == 1

    export_resp = await client.get("/api/admin/users/export", headers=headers, params={"company": "测试公司"})
    assert export_resp.status_code == 200
    account_workbook = load_workbook(BytesIO(export_resp.content))
    account_sheet = account_workbook["账户名单"]
    assert account_sheet.auto_filter.ref == "D1:G2"
    account_values = list(account_sheet.iter_rows(values_only=True))
    assert account_values[0][:6] == ("姓名", "昵称", "手机号", "省份", "公司", "岗位")
    assert account_values[1][0] == "张三"
    assert account_values[1][4] == "测试公司"

    update_resp = await client.put(
        f"/api/admin/users/{created['id']}",
        headers=headers,
        json={
            "phone": "13800001112",
            "login_username": "tech_001",
            "login_password": "",
            "real_name": "李四",
            "nickname": "李四",
            "province": "上海",
            "company": "测试公司",
            "job_role": "技术",
            "profile_verified": True,
        },
    )

    assert update_resp.status_code == 200
    updated = update_resp.json()["data"]
    assert updated["real_name"] == "李四"
    assert updated["job_role"] == "技术"
    assert updated["job_role_label"] == "技术"

    delete_resp = await client.delete(f"/api/admin/users/{created['id']}", headers=headers)
    assert delete_resp.status_code == 200
    assert delete_resp.json()["data"]["deleted"] is True


@pytest.mark.asyncio
async def test_admin_users_list_orders_by_created_at_desc(client, admin_credentials, test_db):
    token = await login_admin(client, admin_credentials)
    headers = admin_headers(token)

    async with test_db() as session:
        old_user = User(
            openid="admin_order_old",
            phone="13800001221",
            real_name="A旧账号",
            nickname="A旧账号",
            province="上海",
            company="排序测试公司",
            job_role="销售",
            total_score=0,
            correct_count=0,
            total_count=0,
            created_at=datetime(2026, 5, 1, 9, 0, 0),
        )
        new_user = User(
            openid="admin_order_new",
            phone="13800001222",
            real_name="Z新账号",
            nickname="Z新账号",
            province="上海",
            company="排序测试公司",
            job_role="销售",
            total_score=0,
            correct_count=0,
            total_count=0,
            created_at=datetime(2026, 5, 2, 9, 0, 0),
        )
        session.add_all([old_user, new_user])
        await session.commit()

    list_resp = await client.get("/api/admin/users", headers=headers, params={"company": "排序测试公司"})
    assert list_resp.status_code == 200
    rows = list_resp.json()["data"]["items"]
    assert [row["real_name"] for row in rows] == ["Z新账号", "A旧账号"]


@pytest.mark.asyncio
async def test_admin_energy_stats_show_earned_energy_without_redemption_deduction(
    client,
    admin_credentials,
    test_db,
):
    token = await login_admin(client, admin_credentials)
    headers = admin_headers(token)

    async with test_db() as session:
        user = User(
            openid="admin_energy_stats_user",
            phone="13800004441",
            real_name="能量学员",
            nickname="能量学员",
            province="上海",
            company="能量公司",
            job_role="销售",
            total_score=500,
            correct_count=50,
            total_count=60,
        )
        other = User(
            openid="admin_energy_stats_other",
            phone="13800004442",
            real_name="低能量学员",
            nickname="低能量学员",
            province="上海",
            company="能量公司",
            job_role="技术",
            total_score=120,
            correct_count=12,
            total_count=20,
        )
        internal = User(
            openid="admin_energy_stats_internal",
            phone="13800004443",
            real_name="内部学员",
            nickname="内部学员",
            province="施耐德电气",
            company="施耐德电气",
            job_role="施耐德内部",
            total_score=999,
            correct_count=99,
            total_count=100,
        )
        session.add_all([user, other, internal])
        await session.flush()
        session.add(
            EnergyRedemptionRecord(
                user_id=user.id,
                client_record_id="energy_stats_redeem_1",
                batch_id="energy_stats_batch",
                product_id="gift",
                product_name="礼品",
                cost=200,
                quantity=1,
                unit_cost=200,
                total_cost=200,
                status="delivered",
            )
        )
        await session.commit()

    resp = await client.get(
        "/api/admin/energy-stats",
        headers=headers,
        params={"company": "能量公司", "limit": 20, "offset": 0},
    )
    assert resp.status_code == 200
    data = resp.json()["data"]
    assert data["total"] == 2
    assert data["items"][0]["name"] == "能量学员"
    assert data["items"][0]["total_score"] == 500
    assert data["items"][0]["correct_count"] == 50
    assert data["items"][0]["total_count"] == 60

    export_resp = await client.get("/api/admin/energy-stats/export", headers=headers)
    assert export_resp.status_code == 200
    workbook = load_workbook(BytesIO(export_resp.content))
    sheet = workbook.active
    exported_rows = list(sheet.iter_rows(min_row=2, values_only=True))
    assert [row[1] for row in exported_rows] == ["能量学员", "低能量学员"]
    assert sum(row[6] for row in exported_rows) == 620
    assert "内部学员" not in {row[1] for row in exported_rows}


@pytest.mark.asyncio
async def test_admin_can_filter_reward_records_by_month(
    client,
    admin_credentials,
    test_db,
):
    token = await login_admin(client, admin_credentials)
    headers = admin_headers(token)

    async with test_db() as session:
        monthly_user = User(
            openid="admin_reward_monthly",
            phone="13800005551",
            real_name="月榜中奖人",
            nickname="月榜中奖人",
            province="上海",
            company="奖励公司",
            job_role="销售",
            total_score=30,
            correct_count=10,
            total_count=10,
        )
        lottery_user = User(
            openid="admin_reward_lottery",
            phone="13800005552",
            real_name="抽奖中奖人",
            nickname="抽奖中奖人",
            province="上海",
            company="奖励公司",
            job_role="技术",
            total_score=20,
            correct_count=8,
            total_count=10,
        )
        other_month_user = User(
            openid="admin_reward_other_month",
            phone="13800005553",
            real_name="其他月份中奖人",
            nickname="其他月份中奖人",
            province="上海",
            company="奖励公司",
            job_role="商务",
            total_score=10,
            correct_count=5,
            total_count=10,
        )
        session.add_all([monthly_user, lottery_user, other_month_user])
        await session.flush()

        session.add(
            MonthlyRankSnapshot(
                month_key="2026-05",
                user_id=monthly_user.id,
                rank=1,
                nickname=monthly_user.nickname,
                real_name=monthly_user.real_name,
                company=monthly_user.company,
                province=monthly_user.province,
                monthly_correct_count=10,
                monthly_total_count=10,
                monthly_time_spent=100,
                reward_amount=30,
                reward_status="issued",
            )
        )
        monthly_tx = EnergyTransaction(
            user_id=monthly_user.id,
            amount=30,
            type="monthly_rank_reward",
            title="2026-05月榜奖励",
            description="月榜第1名，奖励30格施能量",
            related_type="monthly_rank_snapshot",
            related_id=f"2026-05:{monthly_user.id}",
            related_month="2026-05",
            status="issued",
        )
        lottery_tx = EnergyTransaction(
            user_id=lottery_user.id,
            amount=20,
            type="lottery_reward",
            title="2026年05月幸运抽奖",
            description="2026年05月幸运抽奖二等奖，奖励20格施能量",
            related_type="lottery_winner",
            related_id=f"2026-05:{lottery_user.id}",
            related_month="2026-05",
            status="issued",
        )
        other_tx = EnergyTransaction(
            user_id=other_month_user.id,
            amount=10,
            type="monthly_rank_reward",
            title="2026-04月榜奖励",
            description="月榜奖励",
            related_type="monthly_rank_snapshot",
            related_id=f"2026-04:{other_month_user.id}",
            related_month="2026-04",
            status="issued",
        )
        session.add_all([monthly_tx, lottery_tx, other_tx])
        await session.flush()

        draw = LotteryDraw(
            month_key="2026-05",
            participant_month="2026-04",
            status="completed",
            eligible_count=50,
            winner_count=1,
        )
        session.add(draw)
        await session.flush()
        session.add(
            LotteryWinner(
                draw_id=draw.id,
                month_key="2026-05",
                user_id=lottery_user.id,
                prize_level="second",
                prize_name="二等奖",
                reward_amount=20,
                winner_order=1,
                energy_transaction_id=lottery_tx.id,
            )
        )
        await session.commit()

    resp = await client.get(
        "/api/admin/reward-records",
        headers=headers,
        params={"month": "2026-05"},
    )
    assert resp.status_code == 200
    data = resp.json()["data"]
    assert data["total"] == 2
    assert data["monthly_count"] == 1
    assert data["lottery_count"] == 1
    assert data["total_amount"] == 50
    names = {item["name"] for item in data["items"]}
    assert names == {"月榜中奖人", "抽奖中奖人"}
    monthly_item = next(item for item in data["items"] if item["type"] == "monthly_rank_reward")
    lottery_item = next(item for item in data["items"] if item["type"] == "lottery_reward")
    assert monthly_item["rank"] == 1
    assert lottery_item["prize_name"] == "二等奖"

    lottery_resp = await client.get(
        "/api/admin/reward-records",
        headers=headers,
        params={"month": "2026-05", "reward_type": "lottery_reward"},
    )
    assert lottery_resp.status_code == 200
    assert lottery_resp.json()["data"]["total"] == 1
    assert lottery_resp.json()["data"]["items"][0]["name"] == "抽奖中奖人"


@pytest.mark.asyncio
async def test_admin_can_export_all_reward_records_xlsx(
    client,
    admin_credentials,
    test_db,
):
    token = await login_admin(client, admin_credentials)
    headers = admin_headers(token)

    async with test_db() as session:
        monthly_user = User(
            openid="admin_reward_export_monthly",
            phone="13800005601",
            real_name="Reward Export Monthly",
            nickname="Reward Export Monthly",
            province="Shanghai",
            company="Reward Export Co",
            job_role="sales",
            total_score=30,
            correct_count=10,
            total_count=10,
        )
        lottery_user = User(
            openid="admin_reward_export_lottery",
            phone="13800005602",
            real_name="Reward Export Lottery",
            nickname="Reward Export Lottery",
            province="Shanghai",
            company="Reward Export Co",
            job_role="tech",
            total_score=20,
            correct_count=8,
            total_count=10,
        )
        old_user = User(
            openid="admin_reward_export_old",
            phone="13800005603",
            real_name="Reward Export Old",
            nickname="Reward Export Old",
            province="Shanghai",
            company="Reward Export Co",
            job_role="business",
            total_score=10,
            correct_count=5,
            total_count=10,
        )
        session.add_all([monthly_user, lottery_user, old_user])
        await session.flush()

        session.add(
            MonthlyRankSnapshot(
                month_key="2026-05",
                user_id=monthly_user.id,
                rank=1,
                nickname=monthly_user.nickname,
                real_name=monthly_user.real_name,
                company=monthly_user.company,
                province=monthly_user.province,
                monthly_correct_count=10,
                monthly_total_count=10,
                monthly_time_spent=100,
                reward_amount=30,
                reward_status="issued",
            )
        )
        monthly_tx = EnergyTransaction(
            user_id=monthly_user.id,
            amount=30,
            type="monthly_rank_reward",
            title="May monthly reward",
            description="Rank 1 reward",
            related_type="monthly_rank_snapshot",
            related_id=f"2026-05:{monthly_user.id}",
            related_month="2026-05",
            status="issued",
        )
        lottery_tx = EnergyTransaction(
            user_id=lottery_user.id,
            amount=20,
            type="lottery_reward",
            title="June lottery reward",
            description="Lottery first prize",
            related_type="lottery_winner",
            related_id=f"2026-06:{lottery_user.id}",
            related_month="2026-06",
            status="issued",
        )
        old_tx = EnergyTransaction(
            user_id=old_user.id,
            amount=10,
            type="monthly_rank_reward",
            title="April monthly reward",
            description="Old month reward",
            related_type="monthly_rank_snapshot",
            related_id=f"2026-04:{old_user.id}",
            related_month="2026-04",
            status="issued",
        )
        session.add_all([monthly_tx, lottery_tx, old_tx])
        await session.flush()

        draw = LotteryDraw(
            month_key="2026-06",
            participant_month="2026-05",
            status="completed",
            eligible_count=50,
            winner_count=1,
        )
        session.add(draw)
        await session.flush()
        session.add(
            LotteryWinner(
                draw_id=draw.id,
                month_key="2026-06",
                user_id=lottery_user.id,
                prize_level="first",
                prize_name="First prize",
                reward_amount=20,
                winner_order=1,
                energy_transaction_id=lottery_tx.id,
            )
        )
        await session.commit()

    export_resp = await client.get("/api/admin/reward-records/export", headers=headers)
    assert export_resp.status_code == 200
    workbook = load_workbook(BytesIO(export_resp.content))
    sheet = workbook.active
    assert sheet.auto_filter.ref
    rows = list(sheet.iter_rows(values_only=True))
    assert len(rows) == 4
    assert rows[0][:6] == ("月份", "发放时间", "用户", "手机号", "公司", "岗位")
    assert {row[0] for row in rows[1:]} == {"2026-04", "2026-05", "2026-06"}
    assert {row[2] for row in rows[1:]} == {
        "Reward Export Monthly",
        "Reward Export Lottery",
        "Reward Export Old",
    }


@pytest.mark.asyncio
async def test_admin_reports_company_leaderboard_and_participation(
    client,
    admin_credentials,
    test_questions,
    test_db,
):
    token = await login_admin(client, admin_credentials)
    headers = admin_headers(token)

    async with test_db() as session:
        sales = User(
            openid="admin_report_sales",
            phone="13800002221",
            real_name="销售一",
            nickname="销售一",
            province="上海",
            company="报表公司",
            job_role="销售",
            total_score=0,
            correct_count=0,
            total_count=0,
        )
        tech = User(
            openid="admin_report_tech",
            phone="13800002222",
            real_name="技术一",
            nickname="技术一",
            province="上海",
            company="报表公司",
            job_role="技术",
            total_score=0,
            correct_count=0,
            total_count=0,
        )
        shanghai_city = User(
            openid="admin_report_shanghai_city",
            phone="13800002223",
            real_name="上海市员工",
            nickname="上海市员工",
            province="上海市",
            company="上海市公司",
            job_role="销售",
            total_score=0,
            correct_count=0,
            total_count=0,
        )
        company_as_province = User(
            openid="admin_report_bad_province",
            phone="13800002224",
            real_name="公司名省份",
            nickname="公司名省份",
            province="苏州舍得",
            company="苏州舍得电气有限公司",
            job_role="销售",
            total_score=0,
            correct_count=0,
            total_count=0,
        )
        schneider = User(
            openid="admin_report_schneider",
            phone="13800002225",
            real_name="施耐德员工",
            nickname="施耐德员工",
            province="施耐德电气",
            company="施耐德电气",
            job_role="技术",
            total_score=0,
            correct_count=0,
            total_count=0,
        )
        session.add_all([sales, tech, shanghai_city, company_as_province, schneider])
        await session.flush()
        session.add_all(
            [
                AnswerRecord(
                    user_id=sales.id,
                    question_id=test_questions[0].id,
                    selected_answer="A",
                    is_correct=True,
                    score=1,
                    time_spent=10,
                    source="daily",
                    quiz_date="2026-05-04",
                ),
                AnswerRecord(
                    user_id=sales.id,
                    question_id=test_questions[1].id,
                    selected_answer="A",
                    is_correct=True,
                    score=1,
                    time_spent=20,
                    source="daily",
                    quiz_date="2026-05-11",
                ),
                AnswerRecord(
                    user_id=sales.id,
                    question_id=test_questions[0].id,
                    selected_answer="B",
                    is_correct=False,
                    score=0,
                    time_spent=99,
                    source="bank",
                    quiz_date="2026-05-12",
                ),
                AnswerRecord(
                    user_id=tech.id,
                    question_id=test_questions[0].id,
                    selected_answer="B",
                    is_correct=False,
                    score=0,
                    time_spent=30,
                    source="daily",
                    quiz_date="2026-05-04",
                ),
            ]
        )
        await session.commit()

    companies_resp = await client.get("/api/admin/companies", headers=headers)
    assert companies_resp.status_code == 200
    assert "报表公司" in companies_resp.json()["data"]["items"]

    provinces_resp = await client.get("/api/admin/provinces", headers=headers)
    assert provinces_resp.status_code == 200
    provinces = provinces_resp.json()["data"]["items"]
    assert "施耐德电气" in provinces
    assert "上海" in provinces
    assert "上海市" not in provinces
    assert "苏州舍得" not in provinces

    province_companies_resp = await client.get(
        "/api/admin/companies",
        headers=headers,
        params={"province": "上海"},
    )
    assert province_companies_resp.status_code == 200
    assert province_companies_resp.json()["data"]["items"] == ["上海市公司", "报表公司"]

    schneider_companies_resp = await client.get(
        "/api/admin/companies",
        headers=headers,
        params={"province": "施耐德电气"},
    )
    assert schneider_companies_resp.status_code == 200
    assert schneider_companies_resp.json()["data"]["items"] == ["施耐德电气"]

    schneider_users_resp = await client.get(
        "/api/admin/users",
        headers=headers,
        params={"company": "施耐德电气"},
    )
    assert schneider_users_resp.status_code == 200
    schneider_user = schneider_users_resp.json()["data"]["items"][0]
    assert schneider_user["job_role"] == "施耐德内部"
    assert schneider_user["job_role_label"] == "施耐德内部"

    leaderboard_resp = await client.get(
        "/api/admin/company-leaderboard",
        headers=headers,
        params={"company": "报表公司"},
    )
    assert leaderboard_resp.status_code == 200
    leaderboard = leaderboard_resp.json()["data"]["items"]
    assert leaderboard[0]["name"] == "销售一"
    assert leaderboard[0]["weekly_correct_count"] == 2

    total_leaderboard_resp = await client.get(
        "/api/admin/global-leaderboard",
        headers=headers,
        params={"scope": "total", "page": 1},
    )
    assert total_leaderboard_resp.status_code == 200
    total_leaderboard_data = total_leaderboard_resp.json()["data"]
    assert total_leaderboard_data["scope"] == "total"
    assert total_leaderboard_data["page"] == 1
    assert total_leaderboard_data["page_size"] == 20
    total_leaderboard = total_leaderboard_data["items"]
    assert total_leaderboard[0]["name"] == "销售一"
    assert total_leaderboard[0]["total_score"] == 0
    assert total_leaderboard[0]["correct_count"] == 2
    assert total_leaderboard[0]["total_count"] == 2
    assert total_leaderboard[0]["total_time_spent"] == 30

    searched_total_leaderboard_resp = await client.get(
        "/api/admin/global-leaderboard",
        headers=headers,
        params={"scope": "total", "page": 1, "q": "技术一"},
    )
    assert searched_total_leaderboard_resp.status_code == 200
    searched_total_data = searched_total_leaderboard_resp.json()["data"]
    assert searched_total_data["total"] == 1
    assert searched_total_data["items"][0]["name"] == "技术一"

    month_leaderboard_resp = await client.get(
        "/api/admin/global-leaderboard",
        headers=headers,
        params={"scope": "month", "month": "2026-05", "page": 1},
    )
    assert month_leaderboard_resp.status_code == 200
    month_leaderboard_data = month_leaderboard_resp.json()["data"]
    assert month_leaderboard_data["scope"] == "month"
    assert month_leaderboard_data["month"] == "2026-05"
    month_leaderboard = month_leaderboard_data["items"]
    assert month_leaderboard[0]["name"] == "销售一"
    assert month_leaderboard[0]["monthly_correct_count"] == 2
    assert month_leaderboard[0]["monthly_total_count"] == 2

    weekly_resp = await client.get(
        "/api/admin/reports/quiz-participation/weekly",
        headers=headers,
        params={"quiz_date": "2026-05-05"},
    )
    assert weekly_resp.status_code == 200
    weekly_data = weekly_resp.json()["data"]
    assert weekly_data["week_start"] == "2026-05-04"
    assert weekly_data["week_end"] == "2026-05-10"
    assert weekly_data["total_users"] == 2
    assert weekly_data["companies"][0]["company"] == "报表公司"

    searched_weekly_resp = await client.get(
        "/api/admin/reports/quiz-participation/weekly",
        headers=headers,
        params={"quiz_date": "2026-05-05", "q": "技术一"},
    )
    assert searched_weekly_resp.status_code == 200
    searched_weekly_data = searched_weekly_resp.json()["data"]
    assert searched_weekly_data["total_users"] == 1
    assert searched_weekly_data["companies"][0]["participants"][0]["name"] == "技术一"

    weekly_export_resp = await client.get(
        "/api/admin/reports/quiz-participation/weekly/export",
        headers=headers,
        params={"quiz_date": "2026-05-05", "role": "销售,技术"},
    )
    assert weekly_export_resp.status_code == 200
    weekly_workbook = load_workbook(BytesIO(weekly_export_resp.content))
    assert weekly_workbook.sheetnames == ["本周已答题", "本周未答题"]
    weekly_sheet = weekly_workbook["本周已答题"]
    assert weekly_sheet.auto_filter.ref
    assert weekly_workbook["本周未答题"].auto_filter.ref
    weekly_values = list(weekly_sheet.iter_rows(values_only=True))
    weekly_unanswered_values = list(weekly_workbook["本周未答题"].iter_rows(values_only=True))
    assert weekly_values[0][:4] == ("周开始", "周结束", "公司", "员工")
    assert any(row[3] == "销售一" for row in weekly_values)
    assert any(row[3] == "上海市员工" for row in weekly_unanswered_values)

    monthly_resp = await client.get(
        "/api/admin/reports/quiz-participation/monthly",
        headers=headers,
        params={"month": "2026-05", "province": "上海", "company": "报表公司"},
    )
    assert monthly_resp.status_code == 200
    monthly_rows = monthly_resp.json()["data"]["companies"][0]["employees"]
    sales_row = next(row for row in monthly_rows if row["name"] == "销售一")
    tech_row = next(row for row in monthly_rows if row["name"] == "技术一")
    assert sales_row["participated_weeks"] == 2
    assert sales_row["weeks"] == ["2026-05-04", "2026-05-11"]
    assert tech_row["participated_weeks"] == 1

    searched_monthly_resp = await client.get(
        "/api/admin/reports/quiz-participation/monthly",
        headers=headers,
        params={"month": "2026-05", "q": "技术一"},
    )
    assert searched_monthly_resp.status_code == 200
    searched_monthly_rows = searched_monthly_resp.json()["data"]["companies"][0]["employees"]
    assert [row["name"] for row in searched_monthly_rows] == ["技术一"]

    export_resp = await client.get(
        "/api/admin/reports/quiz-participation/monthly/export",
        headers=headers,
        params={"month": "2026-05", "province": "上海", "company": "报表公司"},
    )
    assert export_resp.status_code == 200
    workbook = load_workbook(BytesIO(export_resp.content))
    assert workbook.sheetnames == ["本月已答题", "本月未答题"]
    assert workbook["本月已答题"].auto_filter.ref
    assert workbook["本月未答题"].auto_filter.ref
    answered_values = list(workbook["本月已答题"].iter_rows(values_only=True))
    unanswered_values = list(workbook["本月未答题"].iter_rows(values_only=True))
    assert any(row[1] == "销售一" for row in answered_values)
    assert any(row[1] == "技术一" for row in answered_values)
    assert any(row[1] == "上海市员工" for row in unanswered_values)

    all_export_resp = await client.get(
        "/api/admin/reports/quiz-participation/all/export",
        headers=headers,
    )
    assert all_export_resp.status_code == 200
    all_workbook = load_workbook(BytesIO(all_export_resp.content))
    assert all_workbook.sheetnames == ["全部答题统计"]
    all_sheet = all_workbook["全部答题统计"]
    assert all_sheet.auto_filter.ref
    all_values = list(all_sheet.iter_rows(values_only=True))
    assert all_values[0][:6] == ("公司", "员工", "岗位", "手机号", "参加周数", "答题总数")
    all_sales_row = next(row for row in all_values if row[1] == "销售一")
    assert all_sales_row[4] == 2
    assert all_sales_row[5] == 2
    assert all_sales_row[6] == 2
    assert all_sales_row[8] == 30
    assert all_sales_row[9] == 15


@pytest.mark.asyncio
async def test_admin_global_leaderboard_excludes_users_without_weekly_answers(
    client,
    admin_credentials,
    test_questions,
    test_db,
):
    token = await login_admin(client, admin_credentials)
    headers = admin_headers(token)

    async with test_db() as session:
        session.add_all(
            [
                User(
                    openid="admin_global_fallback_1",
                    phone="13800003331",
                    real_name="累计一",
                    nickname="累计一",
                    province="上海",
                    company="累计公司",
                    job_role="销售",
                    total_score=50,
                    correct_count=5,
                    total_count=7,
                ),
                User(
                    openid="admin_global_fallback_2",
                    phone="13800003332",
                    real_name="累计二",
                    nickname="累计二",
                    province="上海",
                    company="累计公司",
                    job_role="技术",
                    total_score=30,
                    correct_count=3,
                    total_count=6,
                ),
            ]
        )
        await session.flush()
        bank_only_user = (
            await session.execute(select(User).where(User.openid == "admin_global_fallback_1"))
        ).scalar_one()
        session.add(
            AnswerRecord(
                user_id=bank_only_user.id,
                question_id=test_questions[0].id,
                selected_answer="A",
                is_correct=True,
                score=0,
                time_spent=11,
                source="bank",
                quiz_date=None,
            )
        )
        weekly_user = (
            await session.execute(select(User).where(User.openid == "admin_global_fallback_2"))
        ).scalar_one()
        session.add(
            AnswerRecord(
                user_id=weekly_user.id,
                question_id=test_questions[1].id,
                selected_answer="A",
                is_correct=True,
                score=1,
                time_spent=8,
                source="daily",
                quiz_date="2026-05-18",
            )
        )
        await session.commit()

    resp = await client.get(
        "/api/admin/global-leaderboard",
        headers=headers,
        params={"scope": "total", "page": 1},
    )
    assert resp.status_code == 200
    data = resp.json()["data"]
    assert data["total"] == 1
    assert data["items"][0]["name"] == "累计二"
    assert data["items"][0]["correct_count"] == 1
    assert data["items"][0]["total_count"] == 1
    assert data["items"][0]["total_time_spent"] == 8


@pytest.mark.asyncio
async def test_admin_can_list_export_and_update_redemption_orders(
    client,
    admin_credentials,
    test_db,
):
    token = await login_admin(client, admin_credentials)
    headers = admin_headers(token)

    async with test_db() as session:
        user = User(
            openid="admin_order_user",
            phone="13800003333",
            real_name="订单用户",
            nickname="订单用户",
            province="上海",
            company="订单公司",
            job_role="销售",
            total_score=300,
            correct_count=0,
            total_count=0,
        )
        session.add(user)
        await session.flush()
        order = EnergyRedemptionRecord(
            user_id=user.id,
            client_record_id="client-order-1",
            batch_id="batch-1",
            product_id="gift-1",
            product_name="测试礼品",
            cost=50,
            quantity=2,
            unit_cost=50,
            total_cost=100,
            status="pending",
            receiver_name="收货人",
            receiver_phone="13800004444",
            receiver_region="上海市",
            receiver_address="测试地址 1 号",
            receiver_note="尽快发货",
        )
        session.add(order)
        await session.commit()
        order_id = order.id

    list_resp = await client.get(
        "/api/admin/redemption-orders",
        headers=headers,
        params={"company": "订单公司", "q": "测试礼品"},
    )
    assert list_resp.status_code == 200
    data = list_resp.json()["data"]
    assert data["total"] == 1
    row = data["items"][0]
    assert row["product_name"] == "测试礼品"
    assert row["quantity"] == 2
    assert row["receiver_note"] == "尽快发货"
    assert row["status"] == "pending"

    pending_count_resp = await client.get("/api/admin/redemption-orders/pending-count", headers=headers)
    assert pending_count_resp.status_code == 200
    assert pending_count_resp.json()["data"]["pending_count"] == 1

    export_resp = await client.get(
        "/api/admin/redemption-orders/export",
        headers=headers,
        params={"company": "订单公司", "q": "测试礼品"},
    )
    assert export_resp.status_code == 200
    workbook = load_workbook(BytesIO(export_resp.content))
    assert workbook.sheetnames == ["兑换订单"]
    sheet = workbook["兑换订单"]
    assert sheet.auto_filter.ref
    values = list(sheet.iter_rows(values_only=True))
    assert values[0][:5] == ("提交日期", "状态", "用户", "手机号", "公司")
    assert values[1][2] == "订单用户"
    assert values[1][6] == "测试礼品"
    assert values[1][8] == 2
    assert values[1][15] == "尽快发货"

    update_resp = await client.put(
        f"/api/admin/redemption-orders/{order_id}/status",
        headers=headers,
        json={"status": "delivered"},
    )
    assert update_resp.status_code == 200
    assert update_resp.json()["data"]["status"] == "delivered"
    assert update_resp.json()["data"]["status_label"] == "已发货"

    pending_count_after_resp = await client.get("/api/admin/redemption-orders/pending-count", headers=headers)
    assert pending_count_after_resp.status_code == 200
    assert pending_count_after_resp.json()["data"]["pending_count"] == 0

    verify_resp = await client.get(
        "/api/admin/redemption-orders",
        headers=headers,
        params={"status": "delivered"},
    )
    assert verify_resp.status_code == 200
    assert verify_resp.json()["data"]["items"][0]["id"] == order_id
