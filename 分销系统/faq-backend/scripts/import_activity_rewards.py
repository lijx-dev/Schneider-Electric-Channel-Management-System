"""
根据 Excel 获奖名单给指定用户发放施能量，并记录到 energy_transactions 表。

用法:
    cd faq-backend
    # 预览模式（不写入数据库）
    python -m scripts.import_activity_rewards --dry-run

    # 正式执行
    python -m scripts.import_activity_rewards

Excel 文件: 【6月母线豆包互动有奖活动月】获奖名单.xlsx
Sheet: 【6月母线豆包互动有奖活动月】
A 列 = 名字 (real_name)
B 列 = 公司 (company)
E 列 = 奖励施能量（格）(amount)

匹配逻辑:
  - 按 real_name + company 精确匹配 User 表
  - 未匹配到的用户打印警告并跳过
  - 防重复：同一用户 + activity_reward + 2026-06 不可重复发放
"""
import argparse
import asyncio
import os
import sys

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from openpyxl import load_workbook
from sqlalchemy import select

from app.db.session import AsyncSessionLocal, close_db
from app.models.energy import EnergyTransaction
from app.models.user import User

# ========== 配置 ==========
EXCEL_PATH = os.path.join(
    os.path.dirname(__file__), "..", "..",
    "【6月母线豆包互动有奖活动月】获奖名单.xlsx"
)
SHEET_NAME = "【6月母线豆包互动有奖活动月】"
REWARD_MONTH = "2026-06"
REWARD_TYPE = "activity_reward"
REWARD_TITLE = "6月母线豆包互动有奖活动"

COL_NAME = 1      # A 列
COL_COMPANY = 2   # B 列
COL_ENERGY = 5    # E 列
START_ROW = 2     # 数据从第 2 行开始


def read_excel():
    """读取 Excel 文件，返回 [(name, company, energy), ...] 列表，同一用户多次出现时合并金额。"""
    abs_path = os.path.abspath(EXCEL_PATH)
    print(f"读取 Excel: {abs_path}")

    if not os.path.exists(abs_path):
        print(f"错误: 文件不存在 - {abs_path}")
        sys.exit(1)

    wb = load_workbook(abs_path, read_only=True, data_only=True)
    if SHEET_NAME not in wb.sheetnames:
        print(f"错误: 找不到 Sheet '{SHEET_NAME}'，可用: {wb.sheetnames}")
        sys.exit(1)

    ws = wb[SHEET_NAME]
    raw_records = []

    for row in ws.iter_rows(min_row=START_ROW):
        name = row[COL_NAME - 1].value
        company = row[COL_COMPANY - 1].value
        energy = row[COL_ENERGY - 1].value

        if not name or not company:
            continue

        name = str(name).strip()
        company = str(company).strip()
        energy = int(energy) if energy is not None else 0

        if energy <= 0:
            continue

        raw_records.append((name, company, energy))

    wb.close()

    # 合并同一用户（姓名+公司）的多条记录，累加能量值
    merged = {}
    for name, company, energy in raw_records:
        key = (name, company)
        merged[key] = merged.get(key, 0) + energy

    records = [(name, company, total) for (name, company), total in merged.items()]
    records.sort(key=lambda r: r[0])  # 按姓名排序

    if len(raw_records) != len(records):
        duplicates = len(raw_records) - len(records)
        print(f"读取完成: 原始 {len(raw_records)} 条，合并重复用户 {duplicates} 条后共 {len(records)} 条\n")
    else:
        print(f"读取完成: {len(records)} 条有效记录\n")
    return records


async def run(dry_run: bool = False):
    records = read_excel()

    if not records:
        print("没有有效记录，退出")
        return

    # 预览
    print(f"{'名字':<10} {'公司':<30} {'能量':>6}")
    print("-" * 48)
    for name, company, energy in records:
        print(f"{name:<10} {company:<30} {energy:>6}")
    print()

    if dry_run:
        print("=== DRY-RUN 模式，不会写入数据库 ===\n")

    async with AsyncSessionLocal() as session:
        success = 0
        skipped_duplicate = 0
        not_found = 0
        errors = 0

        for name, company, energy in records:
            # 1. 按 real_name + company 匹配用户
            stmt = select(User).where(
                User.real_name == name,
                User.company == company,
            )
            result = await session.execute(stmt)
            user = result.scalar_one_or_none()

            if not user:
                not_found += 1
                print(f"[未匹配] {name} - {company}（系统中无此用户，跳过）")
                continue

            # 2. 防重复检查
            existing_stmt = select(EnergyTransaction).where(
                EnergyTransaction.user_id == user.id,
                EnergyTransaction.type == REWARD_TYPE,
                EnergyTransaction.related_month == REWARD_MONTH,
            )
            existing_result = await session.execute(existing_stmt)
            if existing_result.scalar_one_or_none():
                skipped_duplicate += 1
                print(f"[重复] {name} - {company}（已发放过 {REWARD_MONTH} 活动奖励，跳过）")
                continue

            if dry_run:
                success += 1
                print(f"[预览] {name} - {company} → +{energy} 能量")
                continue

            # 3. 写入数据
            try:
                transaction = EnergyTransaction(
                    user_id=user.id,
                    amount=energy,
                    type=REWARD_TYPE,
                    title=REWARD_TITLE,
                    description=f"6月互动活动奖励 {energy} 格施能量",
                    related_month=REWARD_MONTH,
                    status="issued",
                )
                session.add(transaction)
                user.total_score = (user.total_score or 0) + energy
                success += 1
                print(f"[发放] {name} - {company} → +{energy} 能量")
            except Exception as e:
                errors += 1
                print(f"[错误] {name} - {company}: {e}")

        if not dry_run:
            await session.commit()

    # 输出报告
    print()
    print("=" * 50)
    print("执行报告")
    print("=" * 50)
    print(f"  总记录数:     {len(records)}")
    print(f"  成功发放:     {success}")
    print(f"  重复跳过:     {skipped_duplicate}")
    print(f"  未匹配用户:   {not_found}")
    print(f"  执行错误:     {errors}")
    if dry_run:
        print(f"\n  (DRY-RUN 模式，以上为预览结果，未实际写入数据库)")

    await close_db()


async def main():
    parser = argparse.ArgumentParser(description="导入6月活动奖励")
    parser.add_argument("--dry-run", action="store_true", help="预览模式，不写入数据库")
    args = parser.parse_args()
    await run(dry_run=args.dry_run)


if __name__ == "__main__":
    asyncio.run(main())