"""
从 Excel 表格批量导入分销商用户到数据库

用法:
    cd faq-backend
    python -m scripts.import_users_from_excel

Excel 文件: 副本2025赋能统计表 - 截止12月31日 to rui.xlsm
Sheet 名: 3.2025最新分销商在岗人员名单
A 列 = 公司名称 (company)
C 列 = 姓名 (real_name)
F 列 = 手机号码 (phone)
H 列 = 状态（必须为 Y）

逻辑:
  - 以手机号为唯一标识
  - 如果手机号已存在，则更新 company、real_name、province
  - 如果手机号不存在，则创建新用户（openid 用占位符 "import_手机号"）
  - 省份根据公司名称中的城市关键词自动推断
  - 用户首次在小程序登录时，auth.py 会自动把真实 openid 绑定到该记录
"""
import asyncio
import os
import sys

# 将项目根目录加入 sys.path，使 app 包可被导入
sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from openpyxl import load_workbook
from sqlalchemy import select

from app.db.session import AsyncSessionLocal, init_db
from app.models.user import User


# ========== 配置 ==========
EXCEL_PATH = os.path.join(
    os.path.dirname(__file__), "..",  "..",
    "副本2025赋能统计表 - 截止12月31日 to rui.xlsm"
)
SHEET_NAME = "3.2025最新分销商在岗人员名单"
COL_COMPANY = 1   # A 列
COL_NAME = 3      # C 列
COL_PHONE = 6     # F 列
COL_STATUS = 8    # H 列（必须为 Y 才导入）
START_ROW = 2      # 从第 2 行开始（第 1 行通常是表头）


# ========== 城市 → 省份映射（用于从公司名推断省份）==========
# 键为公司名中可能出现的城市/地区关键词，值为对应省份
CITY_TO_PROVINCE = {
    # 直辖市
    "北京": "北京市", "京": "北京市",
    "上海": "上海市", "沪": "上海市",
    "天津": "天津市", "津": "天津市",
    "重庆": "重庆市", "渝": "重庆市",
    # 广东省
    "广东": "广东省", "广州": "广东省", "深圳": "广东省", "东莞": "广东省",
    "佛山": "广东省", "珠海": "广东省", "惠州": "广东省", "中山": "广东省",
    "江门": "广东省", "汕头": "广东省", "湛江": "广东省", "肇庆": "广东省",
    "茂名": "广东省", "揭阳": "广东省", "梅州": "广东省", "清远": "广东省",
    "阳江": "广东省", "韶关": "广东省", "河源": "广东省", "云浮": "广东省",
    "汕尾": "广东省", "潮州": "广东省",
    # 浙江省
    "浙江": "浙江省", "杭州": "浙江省", "宁波": "浙江省", "温州": "浙江省",
    "嘉兴": "浙江省", "绍兴": "浙江省", "金华": "浙江省", "台州": "浙江省",
    "湖州": "浙江省", "丽水": "浙江省", "衢州": "浙江省", "舟山": "浙江省",
    "义乌": "浙江省",
    # 江苏省
    "江苏": "江苏省", "南京": "江苏省", "苏州": "江苏省", "无锡": "江苏省",
    "常州": "江苏省", "南通": "江苏省", "徐州": "江苏省", "盐城": "江苏省",
    "扬州": "江苏省", "泰州": "江苏省", "镇江": "江苏省", "淮安": "江苏省",
    "连云港": "江苏省", "宿迁": "江苏省", "昆山": "江苏省",
    # 山东省
    "山东": "山东省", "济南": "山东省", "青岛": "山东省", "烟台": "山东省",
    "潍坊": "山东省", "临沂": "山东省", "济宁": "山东省", "淄博": "山东省",
    "威海": "山东省", "德州": "山东省", "聊城": "山东省", "菏泽": "山东省",
    "泰安": "山东省", "枣庄": "山东省", "日照": "山东省", "滨州": "山东省",
    # 四川省
    "四川": "四川省", "成都": "四川省", "绵阳": "四川省", "德阳": "四川省",
    "宜宾": "四川省", "南充": "四川省", "泸州": "四川省", "乐山": "四川省",
    "达州": "四川省", "内江": "四川省", "自贡": "四川省", "遂宁": "四川省",
    # 湖北省
    "湖北": "湖北省", "武汉": "湖北省", "宜昌": "湖北省", "襄阳": "湖北省",
    "荆州": "湖北省", "黄冈": "湖北省", "十堰": "湖北省", "孝感": "湖北省",
    "荆门": "湖北省", "鄂州": "湖北省", "黄石": "湖北省", "咸宁": "湖北省",
    # 湖南省
    "湖南": "湖南省", "长沙": "湖南省", "株洲": "湖南省", "岳阳": "湖南省",
    "衡阳": "湖南省", "常德": "湖南省", "郴州": "湖南省", "邵阳": "湖南省",
    "益阳": "湖南省", "娄底": "湖南省", "永州": "湖南省", "怀化": "湖南省",
    "湘潭": "湖南省", "张家界": "湖南省",
    # 河南省
    "河南": "河南省", "郑州": "河南省", "洛阳": "河南省", "南阳": "河南省",
    "许昌": "河南省", "周口": "河南省", "新乡": "河南省", "商丘": "河南省",
    "信阳": "河南省", "驻马店": "河南省", "焦作": "河南省", "平顶山": "河南省",
    "安阳": "河南省", "开封": "河南省", "漯河": "河南省", "濮阳": "河南省",
    # 河北省
    "河北": "河北省", "石家庄": "河北省", "唐山": "河北省", "保定": "河北省",
    "邯郸": "河北省", "廊坊": "河北省", "沧州": "河北省", "邢台": "河北省",
    "衡水": "河北省", "张家口": "河北省", "承德": "河北省", "秦皇岛": "河北省",
    # 福建省
    "福建": "福建省", "福州": "福建省", "厦门": "福建省", "泉州": "福建省",
    "漳州": "福建省", "莆田": "福建省", "龙岩": "福建省", "三明": "福建省",
    "宁德": "福建省", "南平": "福建省",
    # 安徽省
    "安徽": "安徽省", "合肥": "安徽省", "芜湖": "安徽省", "蚌埠": "安徽省",
    "阜阳": "安徽省", "淮南": "安徽省", "安庆": "安徽省", "马鞍山": "安徽省",
    "滁州": "安徽省", "六安": "安徽省", "宣城": "安徽省", "铜陵": "安徽省",
    # 辽宁省
    "辽宁": "辽宁省", "沈阳": "辽宁省", "大连": "辽宁省", "鞍山": "辽宁省",
    "抚顺": "辽宁省", "锦州": "辽宁省", "营口": "辽宁省", "丹东": "辽宁省",
    "盘锦": "辽宁省", "葫芦岛": "辽宁省", "铁岭": "辽宁省", "朝阳": "辽宁省",
    # 陕西省
    "陕西": "陕西省", "西安": "陕西省", "咸阳": "陕西省", "宝鸡": "陕西省",
    "渭南": "陕西省", "汉中": "陕西省", "榆林": "陕西省", "延安": "陕西省",
    "安康": "陕西省",
    # 江西省
    "江西": "江西省", "南昌": "江西省", "赣州": "江西省", "九江": "江西省",
    "宜春": "江西省", "上饶": "江西省", "抚州": "江西省", "吉安": "江西省",
    "景德镇": "江西省", "萍乡": "江西省", "新余": "江西省",
    # 广西壮族自治区
    "广西": "广西壮族自治区", "南宁": "广西壮族自治区", "柳州": "广西壮族自治区",
    "桂林": "广西壮族自治区", "玉林": "广西壮族自治区", "百色": "广西壮族自治区",
    "梧州": "广西壮族自治区", "贵港": "广西壮族自治区", "钦州": "广西壮族自治区",
    "北海": "广西壮族自治区", "河池": "广西壮族自治区",
    # 山西省
    "山西": "山西省", "太原": "山西省", "大同": "山西省", "临汾": "山西省",
    "运城": "山西省", "长治": "山西省", "晋城": "山西省", "忻州": "山西省",
    "吕梁": "山西省", "晋中": "山西省", "阳泉": "山西省",
    # 云南省
    "云南": "云南省", "昆明": "云南省", "曲靖": "云南省", "大理": "云南省",
    "红河": "云南省", "玉溪": "云南省", "楚雄": "云南省", "昭通": "云南省",
    "保山": "云南省", "丽江": "云南省", "普洱": "云南省",
    # 贵州省
    "贵州": "贵州省", "贵阳": "贵州省", "遵义": "贵州省", "六盘水": "贵州省",
    "黔南": "贵州省", "黔东南": "贵州省", "黔西南": "贵州省", "毕节": "贵州省",
    "安顺": "贵州省", "铜仁": "贵州省",
    # 吉林省
    "吉林": "吉林省", "长春": "吉林省", "四平": "吉林省", "延边": "吉林省",
    "通化": "吉林省", "白城": "吉林省", "松原": "吉林省", "白山": "吉林省",
    "辽源": "吉林省",
    # 黑龙江省
    "黑龙江": "黑龙江省", "哈尔滨": "黑龙江省", "齐齐哈尔": "黑龙江省",
    "大庆": "黑龙江省", "牡丹江": "黑龙江省", "佳木斯": "黑龙江省",
    "绥化": "黑龙江省", "鸡西": "黑龙江省", "鹤岗": "黑龙江省",
    # 甘肃省
    "甘肃": "甘肃省", "兰州": "甘肃省", "天水": "甘肃省", "白银": "甘肃省",
    "庆阳": "甘肃省", "平凉": "甘肃省", "酒泉": "甘肃省", "张掖": "甘肃省",
    "武威": "甘肃省", "定西": "甘肃省", "金昌": "甘肃省", "陇南": "甘肃省",
    # 内蒙古自治区
    "内蒙古": "内蒙古自治区", "内蒙": "内蒙古自治区", "呼和浩特": "内蒙古自治区",
    "包头": "内蒙古自治区", "鄂尔多斯": "内蒙古自治区", "呼伦贝尔": "内蒙古自治区",
    "赤峰": "内蒙古自治区", "通辽": "内蒙古自治区", "乌海": "内蒙古自治区",
    # 新疆维吾尔自治区
    "新疆": "新疆维吾尔自治区", "乌鲁木齐": "新疆维吾尔自治区",
    "克拉玛依": "新疆维吾尔自治区", "喀什": "新疆维吾尔自治区",
    "伊犁": "新疆维吾尔自治区", "昌吉": "新疆维吾尔自治区",
    "阿克苏": "新疆维吾尔自治区", "库尔勒": "新疆维吾尔自治区",
    # 海南省
    "海南": "海南省", "海口": "海南省", "三亚": "海南省", "儋州": "海南省",
    # 宁夏回族自治区
    "宁夏": "宁夏回族自治区", "银川": "宁夏回族自治区", "石嘴山": "宁夏回族自治区",
    "吴忠": "宁夏回族自治区", "固原": "宁夏回族自治区", "中卫": "宁夏回族自治区",
    # 青海省
    "青海": "青海省", "西宁": "青海省", "海东": "青海省",
    # 西藏自治区
    "西藏": "西藏自治区", "拉萨": "西藏自治区", "日喀则": "西藏自治区",
    # 港澳台
    "香港": "香港特别行政区", "港": "香港特别行政区",
    "澳门": "澳门特别行政区", "澳": "澳门特别行政区",
    "台湾": "台湾省", "台北": "台湾省",
}

# 按关键词长度降序排列，优先匹配更精确的词（如"石家庄"优先于"石"）
_SORTED_CITIES = sorted(CITY_TO_PROVINCE.keys(), key=len, reverse=True)


def guess_province(company_name: str) -> str:
    """
    根据公司名称中的城市关键词推断所属省份。
    返回省份名称，如果无法推断则返回空字符串。
    """
    if not company_name:
        return ""
    for city in _SORTED_CITIES:
        if city in company_name:
            return CITY_TO_PROVINCE[city]
    return ""


def read_excel():
    """
    读取 Excel 文件，返回 [(company, real_name, phone, province), ...] 列表。
    跳过手机号为空或 H 列不为 Y 的行。
    """
    abs_path = os.path.abspath(EXCEL_PATH)
    print(f"📂 正在读取 Excel: {abs_path}")

    if not os.path.exists(abs_path):
        print(f"❌ 文件不存在: {abs_path}")
        sys.exit(1)

    wb = load_workbook(abs_path, read_only=True, data_only=True)

    if SHEET_NAME not in wb.sheetnames:
        print(f"❌ 找不到 Sheet: {SHEET_NAME}")
        print(f"   可用的 Sheet: {wb.sheetnames}")
        sys.exit(1)

    ws = wb[SHEET_NAME]
    records = []
    skipped = 0
    no_province = []

    for row in ws.iter_rows(min_row=START_ROW):
        company = row[COL_COMPANY - 1].value
        real_name = row[COL_NAME - 1].value
        phone = row[COL_PHONE - 1].value
        status = row[COL_STATUS - 1].value if len(row) >= COL_STATUS else None

        # H 列必须为 Y 才导入
        if not status or str(status).strip().upper() != "Y":
            skipped += 1
            continue

        # 手机号转字符串并清理
        if phone is not None:
            phone = str(phone).strip()
        if not phone:
            skipped += 1
            continue

        company = str(company).strip() if company else ""
        real_name = str(real_name).strip() if real_name else ""

        # 自动推断省份
        province = guess_province(company)
        if not province and company:
            no_province.append(company)

        records.append((company, real_name, phone, province))

    wb.close()
    print(f"✅ 读取完成: 有效记录 {len(records)} 条, 跳过 {skipped} 条")

    if no_province:
        unique_no_prov = sorted(set(no_province))
        print(f"\n⚠️  以下 {len(unique_no_prov)} 家公司无法自动识别省份（将留空，需手动补充）:")
        for c in unique_no_prov:
            print(f"   {c}")

    return records


async def import_to_db(records):
    """
    将记录写入数据库。
    - 手机号已存在 → 更新 company、real_name、province
    - 手机号不存在 → 创建新用户
    """
    # 确保表结构已创建
    await init_db()

    created = 0
    updated = 0
    errors = []

    async with AsyncSessionLocal() as session:
        for i, (company, real_name, phone, province) in enumerate(records, 1):
            try:
                # 按手机号查找已有用户
                stmt = select(User).where(User.phone == phone)
                result = await session.execute(stmt)
                user = result.scalar_one_or_none()

                if user:
                    # 更新已有用户的公司名、姓名和省份
                    changed = False
                    if company and user.company != company:
                        user.company = company
                        changed = True
                    if real_name and user.real_name != real_name:
                        user.real_name = real_name
                        changed = True
                    if province and user.province != province:
                        user.province = province
                        changed = True
                    if changed:
                        updated += 1
                else:
                    # 创建新用户，openid 使用占位符
                    user = User(
                        openid=f"import_{phone}",
                        phone=phone,
                        real_name=real_name,
                        company=company,
                        province=province,
                        nickname=real_name or f"用户{phone[-4:]}",
                    )
                    session.add(user)
                    created += 1

            except Exception as e:
                errors.append((i, phone, str(e)))

        await session.commit()

    # 输出结果
    print("\n" + "=" * 50)
    print("📊 导入结果汇总")
    print("=" * 50)
    print(f"   新建用户: {created} 条")
    print(f"   更新用户: {updated} 条")
    print(f"   失败记录: {len(errors)} 条")

    if errors:
        print("\n⚠️  失败详情:")
        for row_num, phone, err in errors:
            print(f"   第 {row_num} 条 (手机号: {phone}): {err}")


async def main():
    print("=" * 50)
    print("🚀 分销商用户批量导入工具")
    print("=" * 50)

    records = read_excel()

    if not records:
        print("⚠️  没有有效记录可导入")
        return

    # 预览前 5 条
    print(f"\n📋 数据预览 (前 5 条):")
    print(f"   {'公司名称':<20} {'姓名':<10} {'手机号':<15} {'省份':<10}")
    print(f"   {'-' * 20} {'-' * 10} {'-' * 15} {'-' * 10}")
    for company, name, phone, province in records[:5]:
        print(f"   {company:<20} {name:<10} {phone:<15} {province or '❓未知':<10}")
    if len(records) > 5:
        print(f"   ... 共 {len(records)} 条")

    # 确认导入
    confirm = input("\n确认导入? (y/n): ").strip().lower()
    if confirm != "y":
        print("❌ 已取消导入")
        return

    await import_to_db(records)
    print("\n✅ 导入完成!")


if __name__ == "__main__":
    asyncio.run(main())
