"""Import certificate hierarchy from Excel workbooks.

Usage:
  python scripts/import_siemens_cqc_certificates.py "C:/path/西门子证书查询2026.04.10.xlsx"
  python scripts/import_siemens_cqc_certificates.py "C:/path/西门子证书查询2026.04.10.xlsx" --mode kema
  python scripts/import_siemens_cqc_certificates.py "C:/path/ABB威腾证书-汇总-20260408.xlsx" --mode abb-cqc
  python scripts/import_siemens_cqc_certificates.py "C:/path/ABB威腾证书-汇总-20260408.xlsx" --mode abb-kema
  python scripts/import_siemens_cqc_certificates.py "C:/path/ABB威腾证书-汇总-20260408.xlsx" --mode abb-cb
  python scripts/import_siemens_cqc_certificates.py "C:/path/SE IEC证书.xlsx" --mode se-asta
  python scripts/import_siemens_cqc_certificates.py "C:/path/SE IEC证书.xlsx" --mode se-kema
  python scripts/import_siemens_cqc_certificates.py "C:/path/威腾证书-汇总 20260408.xlsx" --mode weiteng-cqc
  python scripts/import_siemens_cqc_certificates.py "C:/path/威腾证书-汇总 20260408.xlsx" --mode weiteng-kema
  python scripts/import_siemens_cqc_certificates.py "C:/path/证书-Eaton总览-2026.04.08.xlsx" --mode eaton-cqc
  python scripts/import_siemens_cqc_certificates.py "C:/path/证书-Eaton总览-2026.04.08.xlsx" --mode eaton-kema
"""

from __future__ import annotations

import argparse
import asyncio
from datetime import date, datetime
from pathlib import Path
import sys
from typing import Any

import openpyxl
from sqlalchemy import delete

ROOT_DIR = Path(__file__).resolve().parents[1]
if str(ROOT_DIR) not in sys.path:
    sys.path.insert(0, str(ROOT_DIR))

from app.db.session import AsyncSessionLocal, close_db, engine  # noqa: E402
from app.models.certificate import CertificateRecord  # noqa: E402


DEFAULT_SHEET = "CQC汇总2026"
DEFAULT_COMPANY_KEY = "siemens"
DEFAULT_COMPANY_NAME = "西门子"
DEFAULT_CERT_TYPE = "CQC"

SHEET_PROFILES: dict[str, dict[str, Any]] = {
    "cqc": {
        "sheet": "CQC汇总2026",
        "cert_type": "CQC",
        "company_key": "siemens",
        "company_name": "西门子",
        "category_headers": {"产品类型"},
        "model_headers": {"型号"},
        "detail_min_col": 4,
        "header_translations": {},
        "category_translations": {},
    },
    "kema": {
        "sheet": "KEMA-KEUR-2",
        "cert_type": "KEMA",
        "company_key": "siemens",
        "company_name": "西门子",
        "category_headers": {"Product"},
        "model_headers": {"Model"},
        "detail_min_col": 1,
        "header_translations": {
            "Certificate": "证书编号",
            "Product": "产品类型",
            "Model": "型号",
            "Rating": "In（A）",
            "Ue": "额定工作电压 Ue",
            "Uimp": "额定冲击耐受电压 Uimp",
            "Icw (kA)": "Icw(kA)",
            "System": "系统",
            "IP": "IP",
            "Datae of Issue": "初始获证时间",
            "Date of Issue": "初始获证时间",
            "Cert Type": "证书类型",
        },
        "category_translations": {
            "BUSBAR TRUNKING SYSTEMS (BUSWAY)": "母线槽系统（Busway）",
            "Busbar Trunking System (aluminum busbar)": "铝母线槽系统",
            "low-voltage Copper Busbar Trunking System": "低压铜母线槽系统",
            "Low-voltage Cast Resin Busway": "低压浇注树脂母线槽",
            "low-voltage Copper Busbar Trunking System without Tap-off Unit Outlet": "低压铜母线槽系统（无插接箱接口）",
            "Low-voltage Copper Busbar Trunking System without Tap-off Unit Outlet": "低压铜母线槽系统（无插接箱接口）",
            "low-voltage Copper Blow-voltage Copper Busbar Trunking System with/without Tap-off Unit Outlet": "低压铜母线槽系统（带/不带插接箱接口）",
        },
    },
    "abb-cqc": {
        "sheet": "3C汇总",
        "cert_type": "CQC",
        "company_key": "abb",
        "company_name": "ABB威腾",
        "category_headers": {"产品类型"},
        "model_headers": {"型号"},
        "detail_min_col": 3,
        "header_translations": {},
        "category_translations": {},
    },
    "abb-kema": {
        "sheet": "IEC汇总",
        "cert_type": "KEMA",
        "company_key": "abb",
        "company_name": "ABB威腾",
        "category_headers": {"Product"},
        "model_headers": {"Model"},
        "detail_min_col": 1,
        "fallback_headers_by_col": {
            12: "备注",
        },
        "header_translations": {
            "Certificate": "证书编号",
            "Product": "产品类型",
            "Model": "型号",
            "Rating": "In（A）",
            "Ue": "额定工作电压 Ue",
            "Uimp": "额定冲击耐受电压 Uimp",
            "Icw (kA)": "Icw(kA)",
            "System": "系统",
            "IP": "IP",
            "Datae of Issue": "初始获证时间",
            "Date of Issue": "初始获证时间",
            "Cert Type": "证书类型",
        },
        "category_translations": {
            "Busbar Trunking System": "母线槽系统",
            "Low-voltage busbar trunking system (aluminum bar) with/without tap off unit outlet": "低压铝母线槽系统（带/不带插接箱接口）",
            "Low-voltage busbar trunking system (busways) - aluminum bar": "低压铝母线槽系统",
            "Low-voltage busbar trunking system (copper busbar)": "低压铜母线槽系统",
            "Casting low-voltage busbar trunking system (copper bar) without tap-off unit outlet": "浇注式低压铜母线槽系统（无插接箱接口）",
            "Low-voltage Copper Busbar Trunking System with/without Tap-off Unit Outlet": "低压铜母线槽系统（带/不带插接箱接口）",
        },
    },
    "abb-cb": {
        "sheet": "CB证书",
        "cert_type": "CB",
        "company_key": "abb",
        "company_name": "ABB威腾",
        "category_headers": set(),
        "model_headers": {"Model"},
        "synthetic_category": "CB证书",
        "detail_min_col": 1,
        "header_translations": {
            "Certificate": "证书编号",
            "Model": "型号",
            "Manufacturer": "制造商",
            "Rating": "额定参数",
            "Standard": "标准",
            "Date of Issue": "初始获证时间",
        },
        "category_translations": {},
    },
    "se-asta": {
        "sheet": "ASTA-DIAMOND",
        "cert_type": "ASTA-DIAMOND",
        "company_key": "schneider",
        "company_name": "施耐德",
        "category_headers": {"Product Name"},
        "model_headers": {"Model Name"},
        "detail_min_col": 1,
        "header_translations": {
            "DescriptionBrand": "产品描述",
            "Name": "企业名称",
            "Model Name": "型号",
            "Product Name": "产品名称",
        },
        "category_translations": {
            "Busbar Trunking System": "母线槽系统",
            "Plug-In Units (Tap-Offs)": "插接箱（Tap-Offs）",
        },
    },
    "se-kema": {
        "sheet": "KEMA-KEUR",
        "cert_type": "KEMA",
        "company_key": "schneider",
        "company_name": "施耐德",
        "category_headers": set(),
        "model_headers": {"产品型号"},
        "synthetic_category": "KEMA-KEUR",
        "detail_min_col": 1,
        "header_translations": {},
        "category_translations": {},
    },
    "weiteng-cqc": {
        "sheet": "3C汇总2026",
        "cert_type": "CQC",
        "company_key": "weiteng",
        "company_name": "威腾",
        "category_headers": {"产品类型"},
        "model_headers": {"型号"},
        "default_category": "密集型母线槽",
        "detail_min_col": 4,
        "header_translations": {},
        "category_translations": {},
    },
    "weiteng-kema": {
        "sheet": "IEC 汇总",
        "cert_type": "KEMA",
        "company_key": "weiteng",
        "company_name": "威腾",
        "category_headers": {"Product"},
        "model_headers": {"Model"},
        "detail_min_col": 1,
        "header_translations": {
            "Certificate": "证书编号",
            "Product": "产品类型",
            "Model": "型号",
            "Rating": "In（A）",
            "Ue": "额定工作电压 Ue",
            "Uimp": "额定冲击耐受电压 Uimp",
            "Icw (kA)": "Icw(kA)",
            "System": "系统",
            "IP": "IP",
            "Datae of Issue": "初始获证时间",
            "Date of Issue": "初始获证时间",
            "Cert Type": "证书类型",
        },
        "category_translations": {
            "Busway": "母线槽",
            "BUSBAR TRUNKING SYSTEMS (BUSWAY)": "母线槽系统（Busway）",
            "Low-voltage busbar trunking system": "低压母线槽系统",
            "Low-voltage busbar trunking system (copper bar)": "低压铜母线槽系统",
            "Low-voltage Aluminum Busbar Trunking System with/without Tap-off Unit Outlet": "低压铝母线槽系统（带/不带插接箱接口）",
            "Tap-off unit for low-voltage busbar trunking system (TOU)": "低压母线槽插接箱（TOU）",
        },
    },
    "eaton-cqc": {
        "sheet": "3C汇总2026",
        "cert_type": "CQC",
        "company_key": "eaton",
        "company_name": "Eaton",
        "category_headers": {"产品类型"},
        "model_headers": {"型号"},
        "detail_min_col": 4,
        "skip_categories": {"环氧树脂浇筑母线槽"},
        "header_translations": {},
        "category_translations": {},
    },
    "eaton-kema": {
        "sheet": "IEC汇总New",
        "cert_type": "KEMA",
        "company_key": "eaton",
        "company_name": "Eaton",
        "category_headers": {"Product"},
        "model_headers": {"Model"},
        "detail_min_col": 1,
        "header_translations": {
            "Certificate": "证书编号",
            "Product": "产品类型",
            "Model": "型号",
            "Rating": "In（A）",
            "Ue": "额定工作电压 Ue",
            "Uimp": "额定冲击耐受电压 Uimp",
            "Icw (kA)": "Icw(kA)",
            "Datae of Issue": "初始获证时间",
            "Date of Issue": "初始获证时间",
            "Comment": "备注",
            "Cert Type": "证书类型",
        },
        "category_translations": {
            "copper bar": "铜母线",
            "aluminium bar": "铝母线",
            "Low-voltage busbar trunking system (busways) - copper bar": "低压铜母线槽系统",
        },
    },
}


def _clean_cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).replace("\t", "").strip()


def _find_header_row(ws, category_headers: set[str], model_headers: set[str]) -> int:
    for row_idx in range(1, ws.max_row + 1):
        row_values = [_clean_cell(ws.cell(row_idx, col_idx).value) for col_idx in range(1, ws.max_column + 1)]
        has_category = not category_headers or any(value in category_headers for value in row_values)
        has_model = any(value in model_headers for value in row_values)
        if has_category and has_model:
            return row_idx
    raise RuntimeError(f"未在 sheet {ws.title} 中找到产品类型/型号表头行")


def _translate_header(header: str, profile: dict[str, Any]) -> str:
    translations = profile.get("header_translations") or {}
    return translations.get(header, header)


def _translate_category(category: str, profile: dict[str, Any]) -> str:
    translations = profile.get("category_translations") or {}
    return translations.get(category, category)


def _profile_for_sheet(sheet_name: str) -> dict[str, Any]:
    for profile in SHEET_PROFILES.values():
        if profile["sheet"] == sheet_name:
            return profile
    return SHEET_PROFILES["cqc"]


def parse_records(
    workbook_path: Path,
    sheet_name: str = DEFAULT_SHEET,
    company_key: str = DEFAULT_COMPANY_KEY,
    company_name: str = DEFAULT_COMPANY_NAME,
    cert_type: str = DEFAULT_CERT_TYPE,
    profile: dict[str, Any] | None = None,
) -> list[CertificateRecord]:
    profile = profile or _profile_for_sheet(sheet_name)
    wb = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)
    if sheet_name not in wb.sheetnames:
        raise RuntimeError(f"未找到 sheet：{sheet_name}，当前文件包含：{', '.join(wb.sheetnames)}")

    ws = wb[sheet_name]
    has_synthetic_category = bool(profile.get("synthetic_category"))
    category_headers = set(profile.get("category_headers") or (set() if has_synthetic_category else {"产品类型"}))
    model_headers = set(profile.get("model_headers") or {"型号"})

    header_row = _find_header_row(ws, category_headers, model_headers)

    product_type_col = None
    model_col = None
    headers: dict[int, str] = {}
    detail_min_col = int(profile.get("detail_min_col") or 1)

    for col_idx in range(1, ws.max_column + 1):
        header = _clean_cell(ws.cell(header_row, col_idx).value)
        if not header:
            header = (profile.get("fallback_headers_by_col") or {}).get(col_idx, "")
        if category_headers and header in category_headers:
            product_type_col = col_idx
        elif header in model_headers:
            model_col = col_idx
        elif col_idx >= detail_min_col and header:
            translated_header = _translate_header(header, profile)
            if translated_header not in {"产品类型", "型号"}:
                headers[col_idx] = translated_header

    if not model_col or (not product_type_col and not has_synthetic_category):
        raise RuntimeError("表头缺少“产品类型”或“型号”列")

    records: list[CertificateRecord] = []
    current_category = str(profile.get("synthetic_category") or profile.get("default_category") or "")
    current_model = ""
    skip_categories = set(profile.get("skip_categories") or set())
    skip_models = set(profile.get("skip_models") or set())

    for row_idx in range(header_row + 1, ws.max_row + 1):
        raw_category = _clean_cell(ws.cell(row_idx, product_type_col).value) if product_type_col else ""
        raw_model = _clean_cell(ws.cell(row_idx, model_col).value)

        if raw_category:
            current_category = raw_category
        if raw_model:
            current_model = raw_model

        if not current_category or not current_model:
            continue
        if current_category in skip_categories or current_model in skip_models:
            continue

        detail: dict[str, str] = {}
        for col_idx, header in headers.items():
            value = _clean_cell(ws.cell(row_idx, col_idx).value)
            if value:
                detail[header] = value

        if not detail:
            continue

        records.append(
            CertificateRecord(
                company_key=company_key,
                company_name=company_name,
                cert_type=cert_type,
                category_name=_translate_category(current_category, profile),
                model=current_model,
                detail_json=detail,
                source_file=workbook_path.name,
                source_sheet=sheet_name,
                source_row=row_idx,
                sort_order=row_idx,
                is_active=True,
            )
        )

    return records


def _selected_profiles(args: argparse.Namespace) -> list[tuple[str, dict[str, Any]]]:
    if args.mode == "all":
        return [("cqc", SHEET_PROFILES["cqc"]), ("kema", SHEET_PROFILES["kema"])]
    return [(args.mode, SHEET_PROFILES[args.mode])]


async def import_records(args: argparse.Namespace) -> None:
    workbook_path = Path(args.workbook).expanduser().resolve()
    if not workbook_path.exists():
        raise FileNotFoundError(f"文件不存在：{workbook_path}")

    if engine.url.drivername.startswith("sqlite") and str(engine.url.database or "") in ("", ":memory:"):
        raise RuntimeError("当前是内存数据库，导入后不会持久保存。请先配置 DB_TYPE=sqlite 或 DATABASE_URL/MySQL。")

    imported_counts: list[str] = []

    async with AsyncSessionLocal() as session:
        for profile_key, profile in _selected_profiles(args):
            sheet_name = profile["sheet"] if args.mode == "all" else args.sheet or profile["sheet"]
            cert_type = profile["cert_type"] if args.mode == "all" else args.cert_type or profile["cert_type"]
            company_key = args.company_key or profile.get("company_key") or DEFAULT_COMPANY_KEY
            company_name = args.company_name or profile.get("company_name") or DEFAULT_COMPANY_NAME
            records = parse_records(
                workbook_path=workbook_path,
                sheet_name=sheet_name,
                company_key=company_key,
                company_name=company_name,
                cert_type=cert_type,
                profile=profile,
            )

            await session.execute(
                delete(CertificateRecord)
                .where(CertificateRecord.company_key == company_key)
                .where(CertificateRecord.cert_type == cert_type)
                .where(CertificateRecord.source_sheet == sheet_name)
            )
            session.add_all(records)
            imported_counts.append(f"{profile_key}:{len(records)}")

        await session.commit()

    await close_db()
    print(f"Imported records from {workbook_path.name}: {', '.join(imported_counts)}")


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Import certificate records.")
    parser.add_argument("workbook", help="Excel workbook path")
    parser.add_argument(
        "--mode",
        choices=[
            "all",
            "cqc",
            "kema",
            "abb-cqc",
            "abb-kema",
            "abb-cb",
            "se-asta",
            "se-kema",
            "weiteng-cqc",
            "weiteng-kema",
            "eaton-cqc",
            "eaton-kema",
        ],
        default="all",
    )
    parser.add_argument("--sheet", default=None, help="Override sheet name for the selected single mode")
    parser.add_argument("--company-key", default=None)
    parser.add_argument("--company-name", default=None)
    parser.add_argument("--cert-type", default=None, help="Override certificate type for the selected single mode")
    return parser


if __name__ == "__main__":
    asyncio.run(import_records(build_parser().parse_args()))
