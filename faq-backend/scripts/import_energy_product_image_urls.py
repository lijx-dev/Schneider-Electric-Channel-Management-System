"""Import image URLs from a filled workbook into the energy product seed file."""

from __future__ import annotations

import argparse
import json
import os
import subprocess
import sys
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[2]
DEFAULT_WORKBOOK = ROOT / "outputs" / "energy-image-url-template" / "能量商城商品图片URL填写模板.xlsx"
SEED_FILE = ROOT / "faq-backend" / "scripts" / "energy_products_seed.json"


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--input",
        default=str(DEFAULT_WORKBOOK),
        help="已填写 image_url 的 Excel 文件路径",
    )
    parser.add_argument(
        "--sheet",
        default="图片URL模板",
        help="商品数据所在工作表名称，默认：图片URL模板",
    )
    parser.add_argument(
        "--write-db",
        action="store_true",
        help="更新种子文件后，顺手执行 seed_energy_products.py 写入数据库",
    )
    parser.add_argument(
        "--clear-missing",
        action="store_true",
        help="如果 Excel 中 image_url 为空，则清空种子中的 image_url；默认空白不覆盖",
    )
    return parser.parse_args()


def load_rows(workbook_path: Path, sheet_name: str) -> dict[str, str | None]:
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    if sheet_name not in workbook.sheetnames:
        raise ValueError(f"工作表不存在：{sheet_name}")

    sheet = workbook[sheet_name]
    header_row = next(sheet.iter_rows(min_row=4, max_row=4, values_only=True))
    header_map = {str(value).strip(): index for index, value in enumerate(header_row) if value is not None}
    required_headers = ["product_id", "image_url（你填写）"]
    missing = [header for header in required_headers if header not in header_map]
    if missing:
        raise ValueError(f"模板缺少必要列：{', '.join(missing)}")

    product_col = header_map["product_id"]
    image_url_col = header_map["image_url（你填写）"]
    result: dict[str, str | None] = {}

    for row in sheet.iter_rows(min_row=5, values_only=True):
        product_id = str(row[product_col] or "").strip().lower()
        if not product_id:
            continue
        raw_url = row[image_url_col]
        image_url = str(raw_url).strip() if raw_url not in (None, "") else None
        result[product_id] = image_url

    return result


def update_seed(image_url_map: dict[str, str | None], clear_missing: bool) -> tuple[int, int]:
    items = json.loads(SEED_FILE.read_text(encoding="utf-8"))
    updated_count = 0
    touched_count = 0

    for item in items:
        product_id = str(item.get("product_id") or "").strip().lower()
        if product_id not in image_url_map:
            continue
        touched_count += 1
        new_url = image_url_map[product_id]
        if new_url is None and not clear_missing:
            continue

        next_value = new_url or ""
        if str(item.get("image_url") or "").strip() != next_value:
            item["image_url"] = next_value
            updated_count += 1

    SEED_FILE.write_text(json.dumps(items, ensure_ascii=False, indent=4) + "\n", encoding="utf-8")
    return touched_count, updated_count


def seed_database() -> None:
    env = os.environ.copy()
    env["PYTHONPATH"] = "."
    subprocess.run(
        [sys.executable, "scripts/seed_energy_products.py"],
        cwd=ROOT / "faq-backend",
        env=env,
        check=True,
    )


def main() -> int:
    args = parse_args()
    workbook_path = Path(args.input).resolve()
    if not workbook_path.exists():
        raise FileNotFoundError(f"Excel 文件不存在：{workbook_path}")

    image_url_map = load_rows(workbook_path, args.sheet)
    touched_count, updated_count = update_seed(image_url_map, clear_missing=args.clear_missing)

    print(f"识别到 {len(image_url_map)} 条商品行。")
    print(f"命中种子商品 {touched_count} 条，实际更新 image_url {updated_count} 条。")
    print(f"已写回种子文件：{SEED_FILE}")

    if args.write_db:
        seed_database()
        print("已执行数据库回写。")

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
