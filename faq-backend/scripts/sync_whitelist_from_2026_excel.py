"""Sync distributor phone whitelist from the 2026 staff workbook.

Usage:
    python -m scripts.sync_whitelist_from_2026_excel --dry-run
    python -m scripts.sync_whitelist_from_2026_excel --apply
"""

from __future__ import annotations

import argparse
import asyncio
import json
import os
import re
import sys
from dataclasses import asdict, dataclass
from datetime import datetime
from decimal import Decimal
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook
from sqlalchemy import delete, select

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from app.db.session import AsyncSessionLocal, close_db, init_db  # noqa: E402
from app.models.user import User  # noqa: E402
from scripts.import_users_from_excel import guess_province  # noqa: E402


DEFAULT_EXCEL_PATH = (
    r"C:\Users\86135\Documents\WeChat Files\wxid_stm8la2kv7jr22"
    r"\FileStorage\File\2026-05\2026分销商在岗人员名单（需持续摸查）.xlsx"
)
DEFAULT_SHEET_NAME = "2026最新分销商在岗人员名单"

COL_COMPANY = 1
COL_NAME = 3
COL_PHONE = 6
COL_STATUS = 8
START_ROW = 2

PHONE_RE = re.compile(r"1\d{10}")


@dataclass(frozen=True)
class WhitelistRecord:
    phone: str
    real_name: str
    company: str
    province: str


def normalize_phone(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    if isinstance(value, Decimal):
        value = int(value) if value == value.to_integral_value() else str(value)
    text = str(value).strip()
    if not text:
        return ""
    digits = re.sub(r"\D", "", text)
    match = PHONE_RE.search(digits)
    return match.group(0) if match else digits


def clean_text(value: object) -> str:
    return str(value).strip() if value is not None else ""


def read_whitelist_records(excel_path: Path, sheet_name: str) -> list[WhitelistRecord]:
    if not excel_path.exists():
        raise FileNotFoundError(f"Excel file not found: {excel_path}")

    wb = load_workbook(excel_path, read_only=True, data_only=True)
    try:
        if sheet_name not in wb.sheetnames:
            raise RuntimeError(f"Sheet {sheet_name!r} not found. Available: {wb.sheetnames}")

        ws = wb[sheet_name]
        records_by_phone: dict[str, tuple[int, WhitelistRecord]] = {}
        duplicate_rows: list[tuple[int, str]] = []
        invalid_phones: list[tuple[int, str]] = []

        for row_index, row in enumerate(ws.iter_rows(min_row=START_ROW), START_ROW):
            company = clean_text(row[COL_COMPANY - 1].value)
            real_name = clean_text(row[COL_NAME - 1].value)
            raw_phone = row[COL_PHONE - 1].value
            phone = normalize_phone(raw_phone)
            status = clean_text(row[COL_STATUS - 1].value).upper() if len(row) >= COL_STATUS else ""

            if not phone:
                continue
            if not PHONE_RE.fullmatch(phone):
                invalid_phones.append((row_index, clean_text(raw_phone)))
                continue

            record = WhitelistRecord(
                phone=phone,
                real_name=real_name,
                company=company,
                province=guess_province(company),
            )
            if phone in records_by_phone:
                duplicate_rows.append((row_index, phone))
            priority = 1 if status == "Y" else 0
            existing_priority = records_by_phone.get(phone, (-1, record))[0]
            if priority >= existing_priority:
                records_by_phone[phone] = (priority, record)

        print(f"Excel sheet: {sheet_name}")
        print(f"Valid phones: {len(records_by_phone)}")
        print(f"Duplicate phone rows collapsed: {len(duplicate_rows)}")
        print(f"Invalid phone rows skipped: {len(invalid_phones)}")
        if duplicate_rows[:10]:
            print("Duplicate sample:", duplicate_rows[:10])
        if invalid_phones[:10]:
            print("Invalid sample:", invalid_phones[:10])

        return sorted((item[1] for item in records_by_phone.values()), key=lambda item: item.phone)
    finally:
        wb.close()


def is_admin_user(user: User) -> bool:
    if user.login_username:
        return True
    admin_tokens = ("管理员", "admin", "administrator")
    values = [user.real_name, user.nickname, user.company, user.login_username, user.openid]
    return any(
        token in str(value or "").lower()
        for value in values
        for token in admin_tokens
    )


def user_to_backup(user: User) -> dict[str, object]:
    return {
        "id": user.id,
        "openid": user.openid,
        "nickname": user.nickname,
        "phone": user.phone,
        "login_username": user.login_username,
        "login_password": user.login_password,
        "real_name": user.real_name,
        "province": user.province,
        "company": user.company,
        "total_score": user.total_score,
        "correct_count": user.correct_count,
        "total_count": user.total_count,
        "profile_verified": user.profile_verified,
        "weak_categories": user.weak_categories,
        "last_study_at": user.last_study_at.isoformat() if user.last_study_at else None,
        "created_at": user.created_at.isoformat() if user.created_at else None,
        "updated_at": user.updated_at.isoformat() if user.updated_at else None,
    }


def summarize_sample(records: Iterable[object], limit: int = 10) -> list[object]:
    return list(records)[:limit]


async def sync_records(records: list[WhitelistRecord], apply: bool) -> None:
    await init_db()
    desired_by_phone = {record.phone: record for record in records}
    desired_phones = set(desired_by_phone)

    async with AsyncSessionLocal() as session:
        existing = (await session.execute(select(User))).scalars().all()
        existing_by_phone: dict[str, User] = {}
        duplicate_existing_phones: list[tuple[str, str, str]] = []
        for user in existing:
            normalized_phone = normalize_phone(user.phone)
            if not normalized_phone:
                continue
            if normalized_phone in existing_by_phone:
                duplicate_existing_phones.append((normalized_phone, existing_by_phone[normalized_phone].id, user.id))
                continue
            existing_by_phone[normalized_phone] = user

        protected_admins = [user for user in existing if is_admin_user(user)]
        to_delete = [
            user
            for user in existing
            if user.phone
            and normalize_phone(user.phone) not in desired_phones
            and not is_admin_user(user)
        ]

        to_create = [
            record for record in records if record.phone not in existing_by_phone
        ]

        to_update: list[tuple[User, WhitelistRecord, dict[str, tuple[object, object]]]] = []
        for phone, record in desired_by_phone.items():
            user = existing_by_phone.get(phone)
            if not user:
                continue
            changes: dict[str, tuple[object, object]] = {}
            for attr in ("real_name", "company", "province"):
                new_value = getattr(record, attr)
                if new_value and getattr(user, attr) != new_value:
                    changes[attr] = (getattr(user, attr), new_value)
            if user.phone != record.phone:
                changes["phone"] = (user.phone, record.phone)
            if changes:
                to_update.append((user, record, changes))

        print("\nDatabase preview")
        print(f"Existing users: {len(existing)}")
        print(f"Protected admin/password users: {len(protected_admins)}")
        print(f"Create: {len(to_create)}")
        print(f"Update: {len(to_update)}")
        print(f"Delete ordinary phone whitelist users: {len(to_delete)}")
        print(f"Duplicate normalized phones in DB ignored: {len(duplicate_existing_phones)}")
        if duplicate_existing_phones[:10]:
            print("DB duplicate sample:", duplicate_existing_phones[:10])
        print("Create sample:", [asdict(item) for item in summarize_sample(to_create, 5)])
        print(
            "Update sample:",
            [
                {"phone": record.phone, "changes": changes}
                for _, record, changes in summarize_sample(to_update, 5)
            ],
        )
        print(
            "Delete sample:",
            [
                {
                    "id": user.id,
                    "phone": user.phone,
                    "real_name": user.real_name,
                    "company": user.company,
                }
                for user in summarize_sample(to_delete, 10)
            ],
        )
        print(
            "Protected admin sample:",
            [
                {
                    "id": user.id,
                    "phone": user.phone,
                    "login_username": user.login_username,
                    "real_name": user.real_name,
                    "company": user.company,
                }
                for user in summarize_sample(protected_admins, 10)
            ],
        )

        if not apply:
            print("\nDry run only. Re-run with --apply to write changes.")
            return

        backup_dir = Path(__file__).resolve().parents[2] / "tmp" / "whitelist_backups"
        backup_dir.mkdir(parents=True, exist_ok=True)
        backup_path = backup_dir / f"users_backup_{datetime.now():%Y%m%d_%H%M%S}.json"
        backup_path.write_text(
            json.dumps([user_to_backup(user) for user in existing], ensure_ascii=False, indent=2),
            encoding="utf-8",
        )
        print(f"\nBackup written: {backup_path}")

        for user, record, _ in to_update:
            user.phone = record.phone
            user.real_name = record.real_name or user.real_name
            user.company = record.company or user.company
            user.province = record.province or user.province

        for record in to_create:
            session.add(
                User(
                    openid=f"import_{record.phone}",
                    phone=record.phone,
                    real_name=record.real_name,
                    company=record.company,
                    province=record.province,
                    nickname=record.real_name or f"用户{record.phone[-4:]}",
                )
            )

        if to_delete:
            await session.execute(delete(User).where(User.id.in_([user.id for user in to_delete])))

        await session.commit()

    async with AsyncSessionLocal() as verify_session:
        users = (await verify_session.execute(select(User))).scalars().all()
        phone_users = [user for user in users if user.phone]
        ordinary_extra = [
            user.phone
            for user in phone_users
            if user.phone not in desired_phones and not is_admin_user(user)
        ]
        missing = [
            phone
            for phone in desired_phones
            if not any(user.phone == phone for user in phone_users)
        ]
        print("\nVerification")
        print(f"Phone users after sync: {len(phone_users)}")
        print(f"Missing desired phones: {len(missing)}")
        print(f"Unexpected ordinary phones remaining: {len(ordinary_extra)}")
        if missing[:10]:
            print("Missing sample:", missing[:10])
        if ordinary_extra[:10]:
            print("Unexpected sample:", ordinary_extra[:10])


async def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--excel", default=DEFAULT_EXCEL_PATH)
    parser.add_argument("--sheet", default=DEFAULT_SHEET_NAME)
    parser.add_argument("--apply", action="store_true")
    parser.add_argument("--dry-run", action="store_true")
    args = parser.parse_args()

    if args.apply and args.dry_run:
        raise SystemExit("Choose either --apply or --dry-run, not both.")

    records = read_whitelist_records(Path(args.excel), args.sheet)
    try:
        await sync_records(records, apply=args.apply)
    finally:
        await close_db()


if __name__ == "__main__":
    asyncio.run(main())
