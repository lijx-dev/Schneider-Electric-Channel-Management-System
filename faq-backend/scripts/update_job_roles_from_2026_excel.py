"""Update users.job_role from the 2026 distributor staff workbook.

Usage:
    python scripts/update_job_roles_from_2026_excel.py --dry-run
    python scripts/update_job_roles_from_2026_excel.py --apply
"""

from __future__ import annotations

import argparse
import asyncio
import os
import re
import sys
from collections import Counter
from dataclasses import dataclass
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import bindparam, text

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from app.db.session import close_db, engine  # noqa: E402


DEFAULT_EXCEL_PATH = r"D:\FAQ源数据集\副本2026分销商在岗人员名单（需持续摸查）(1).xlsx"
DEFAULT_SHEET_NAME = "2026最新分销商在岗人员名单"

COL_COMPANY = 1
COL_NAME = 3
COL_JOB_DESC = 5
COL_PHONE = 6
COL_MINI_PROGRAM = 11
START_ROW = 2

PHONE_RE = re.compile(r"1\d{10}")
SCHNEIDER_COMPANY = "施耐德电气"
SCHNEIDER_INTERNAL_ROLE = "施耐德内部"

@dataclass(frozen=True)
class JobRoleRecord:
    row_index: int
    phone: str
    real_name: str
    company: str
    raw_job_desc: str
    job_role: str


def clean_text(value: object) -> str:
    return str(value).strip() if value is not None else ""


def normalize_phone(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    if isinstance(value, Decimal):
        value = int(value) if value == value.to_integral_value() else str(value)
    digits = re.sub(r"\D", "", str(value).strip())
    match = PHONE_RE.search(digits)
    return match.group(0) if match else digits


def normalize_job_role(job_desc: str, company: str = "") -> str:
    if SCHNEIDER_COMPANY in clean_text(company):
        return SCHNEIDER_INTERNAL_ROLE
    return clean_text(job_desc)[:20]


def read_records(excel_path: Path, sheet_name: str) -> list[JobRoleRecord]:
    if not excel_path.exists():
        raise FileNotFoundError(f"Excel file not found: {excel_path}")

    wb = load_workbook(excel_path, read_only=True, data_only=True)
    try:
        if sheet_name not in wb.sheetnames:
            raise RuntimeError(f"Sheet {sheet_name!r} not found. Available: {wb.sheetnames}")

        ws = wb[sheet_name]
        records_by_phone: dict[str, JobRoleRecord] = {}
        duplicate_rows: list[tuple[int, str]] = []
        invalid_rows: list[tuple[int, str]] = []
        skipped_not_mini_program = 0

        for row_index, row in enumerate(ws.iter_rows(min_row=START_ROW), START_ROW):
            mini_program = clean_text(row[COL_MINI_PROGRAM - 1].value if len(row) >= COL_MINI_PROGRAM else "")
            if "小程序" not in mini_program:
                skipped_not_mini_program += 1
                continue

            phone = normalize_phone(row[COL_PHONE - 1].value if len(row) >= COL_PHONE else "")
            if not PHONE_RE.fullmatch(phone):
                invalid_rows.append((row_index, clean_text(row[COL_PHONE - 1].value if len(row) >= COL_PHONE else "")))
                continue

            company = clean_text(row[COL_COMPANY - 1].value if len(row) >= COL_COMPANY else "")
            raw_job_desc = clean_text(row[COL_JOB_DESC - 1].value if len(row) >= COL_JOB_DESC else "")
            record = JobRoleRecord(
                row_index=row_index,
                phone=phone,
                real_name=clean_text(row[COL_NAME - 1].value if len(row) >= COL_NAME else ""),
                company=company,
                raw_job_desc=raw_job_desc,
                job_role=normalize_job_role(raw_job_desc, company),
            )
            if phone in records_by_phone:
                duplicate_rows.append((row_index, phone))
            records_by_phone[phone] = record

        records = sorted(records_by_phone.values(), key=lambda item: item.phone)
        print("Excel preview")
        print(f"Sheet: {sheet_name}")
        print(f"Eligible mini-program rows: {len(records)}")
        print(f"Skipped non mini-program rows: {skipped_not_mini_program}")
        print(f"Duplicate phones collapsed: {len(duplicate_rows)}")
        print(f"Invalid phone rows skipped: {len(invalid_rows)}")
        print(f"Raw job distribution: {dict(Counter(record.raw_job_desc or '<empty>' for record in records))}")
        print(f"Role distribution: {dict(Counter(record.job_role or '<empty>' for record in records))}")
        if invalid_rows[:10]:
            print("Invalid phone sample:", invalid_rows[:10])
        if duplicate_rows[:10]:
            print("Duplicate phone sample:", duplicate_rows[:10])
        print(
            "Record sample:",
            [
                {
                    "row": record.row_index,
                    "phone": record.phone,
                    "name": record.real_name,
                    "company": record.company,
                    "job": record.raw_job_desc,
                    "role": record.job_role,
                }
                for record in records[:8]
            ],
        )
        return records
    finally:
        wb.close()


async def ensure_job_role_column() -> None:
    async with engine.connect() as conn:
        result = await conn.execute(text("SHOW COLUMNS FROM users LIKE 'job_role'"))
        if result.first() is None:
            raise RuntimeError(
                "Cloud MySQL table users is missing column job_role. "
                "Run alembic upgrade head before applying this update."
            )


async def preview_or_apply(records: list[JobRoleRecord], apply: bool) -> None:
    if not records:
        print("No eligible records found.")
        return

    await ensure_job_role_column()

    desired_by_phone = {record.phone: record for record in records}
    async with engine.begin() as conn:
        result = await conn.execute(
            text(
                """
                SELECT id, phone, real_name, company, job_role
                FROM users
                WHERE phone IN :phones
                """
            ).bindparams(bindparam("phones", expanding=True)),
            {"phones": tuple(desired_by_phone.keys())},
        )
        db_rows = [dict(row._mapping) for row in result]

        matched_by_phone = {normalize_phone(row.get("phone")): row for row in db_rows}
        missing = [record for record in records if record.phone not in matched_by_phone]
        changes: list[tuple[dict[str, object], JobRoleRecord]] = []
        unchanged: list[tuple[dict[str, object], JobRoleRecord]] = []

        for record in records:
            row = matched_by_phone.get(record.phone)
            if not row:
                continue
            current_role = clean_text(row.get("job_role"))
            if current_role != record.job_role:
                changes.append((row, record))
            else:
                unchanged.append((row, record))

        print("\nDatabase preview")
        print(f"Matched users: {len(db_rows)}")
        print(f"Will update: {len(changes)}")
        print(f"Already unchanged: {len(unchanged)}")
        print(f"Missing users by phone: {len(missing)}")
        print(
            "Update sample:",
            [
                {
                    "phone": record.phone,
                    "name": row.get("real_name") or record.real_name,
                    "company": row.get("company") or record.company,
                    "raw_job": record.raw_job_desc,
                    "old_role": row.get("job_role") or "",
                    "new_role": record.job_role,
                }
                for row, record in changes[:10]
            ],
        )
        if missing[:10]:
            print(
                "Missing sample:",
                [
                    {
                        "row": record.row_index,
                        "phone": record.phone,
                        "name": record.real_name,
                        "company": record.company,
                        "job": record.raw_job_desc,
                    }
                    for record in missing[:10]
                ],
            )

        if not apply:
            print("\nDry run only. Re-run with --apply to write users.job_role.")
            return

        for row, record in changes:
            await conn.execute(
                text("UPDATE users SET job_role = :job_role, updated_at = NOW() WHERE id = :user_id"),
                {"job_role": record.job_role, "user_id": row["id"]},
            )
        print(f"\nApplied updates: {len(changes)}")


async def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--excel", default=DEFAULT_EXCEL_PATH)
    parser.add_argument("--sheet", default=DEFAULT_SHEET_NAME)
    parser.add_argument("--apply", action="store_true")
    parser.add_argument("--dry-run", action="store_true")
    args = parser.parse_args()

    if args.apply and args.dry_run:
        raise SystemExit("Choose either --apply or --dry-run, not both.")

    records = read_records(Path(args.excel), args.sheet)
    try:
        await preview_or_apply(records, apply=args.apply)
    finally:
        await close_db()


if __name__ == "__main__":
    asyncio.run(main())
