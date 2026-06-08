"""Import a single-sheet classified question bank XLSX into the questions table.

Expected columns in sheet "题库" (or the active sheet):
1. 题目
2. 正确答案
3. 选项A
4. 选项B
5. 选项C
6. 选项D
7. 选项E
8. 选项F
9. 解析
10. 题型分类

The script is idempotent by question_type + normalized content:
- existing rows are updated in place
- missing rows are inserted

Usage:
  cd faq-backend
  venv\\Scripts\\python scripts/import_classified_question_bank_xlsx.py "C:\\path\\to\\file.xlsx"
  venv\\Scripts\\python scripts/import_classified_question_bank_xlsx.py "C:\\path\\to\\file.xlsx" --apply
"""

from __future__ import annotations

import argparse
import asyncio
import os
import sys
from collections import Counter
from dataclasses import dataclass

from openpyxl import load_workbook
from sqlalchemy import select

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))


DEFAULT_EXCEL_PATH = (
    r"C:\Users\86135\Documents\WeChat Files\wxid_stm8la2kv7jr22\FileStorage\File"
    r"\2026-05\题库-33-48 （已分类）.xlsx"
)

TRUE_FALSE_VALUES = {
    "对": "对",
    "错": "错",
    "正确": "对",
    "错误": "错",
    "是": "对",
    "否": "错",
    "TRUE": "对",
    "FALSE": "错",
}

CHOICE_LETTERS = "ABCDEF"


@dataclass(frozen=True)
class ParsedQuestion:
    question_type: str
    content: str
    options: list[str] | None
    answer: str
    explanation: str
    category: str | None


def clean_text(value) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    if text.endswith(".0"):
        try:
            return str(int(float(text)))
        except ValueError:
            pass
    return text


def normalize_content(value: str) -> str:
    return " ".join(clean_text(value).split())


def normalize_category(value: str) -> str | None:
    text = clean_text(value)
    return text or None


def normalize_choice_answer(answer: str) -> str:
    letters: list[str] = []
    for char in answer.upper():
        if char in CHOICE_LETTERS and char not in letters:
            letters.append(char)
    return "".join(letters)


def detect_question_type(content: str, answer: str, options: list[str]) -> str:
    if options:
        normalized_answer = normalize_choice_answer(answer)
        if answer.upper() in TRUE_FALSE_VALUES:
            return "true_false"
        if len(normalized_answer) > 1:
            return "multiple_choice"
        return "single_choice"

    if any(marker in content for marker in ("（）", "()", "____", "___")):
        return "fill_blank"

    return "short_answer"


def normalize_answer(question_type: str, answer: str) -> str:
    text = clean_text(answer)
    if question_type == "true_false":
        return TRUE_FALSE_VALUES.get(text.upper(), TRUE_FALSE_VALUES.get(text, text))
    if question_type in {"single_choice", "multiple_choice"}:
        normalized = normalize_choice_answer(text)
        return normalized or text.upper()
    return text


def normalize_options(question_type: str, options: list[str]) -> list[str] | None:
    if question_type in {"single_choice", "multiple_choice"}:
        return options
    if question_type == "true_false":
        return options or ["对", "错"]
    return None


def parse_workbook(file_path: str) -> list[ParsedQuestion]:
    workbook = load_workbook(file_path, read_only=True, data_only=True)
    worksheet = workbook["题库"] if "题库" in workbook.sheetnames else workbook.active

    questions: list[ParsedQuestion] = []
    for row in worksheet.iter_rows(min_row=2, values_only=True):
        content = clean_text(row[0] if len(row) > 0 else None)
        if not content:
            continue

        raw_answer = clean_text(row[1] if len(row) > 1 else None)
        options = [
            clean_text(row[col_idx])
            for col_idx in range(2, 8)
            if len(row) > col_idx and clean_text(row[col_idx])
        ]
        explanation = clean_text(row[8] if len(row) > 8 else None)
        category = normalize_category(row[9] if len(row) > 9 else None)

        question_type = detect_question_type(content, raw_answer, options)
        answer = normalize_answer(question_type, raw_answer)
        normalized_options = normalize_options(question_type, options)

        if not answer:
            continue

        questions.append(
            ParsedQuestion(
                question_type=question_type,
                content=content,
                options=normalized_options,
                answer=answer,
                explanation=explanation,
                category=category,
            )
        )

    return questions


async def sync_questions(file_path: str, apply: bool) -> dict:
    from app.db.session import AsyncSessionLocal, engine, init_db
    from app.models.question import Question

    parsed_questions = parse_workbook(file_path)
    parsed_by_key: dict[tuple[str, str], ParsedQuestion] = {}
    duplicate_excel_rows = 0
    for item in parsed_questions:
        key = (item.question_type, normalize_content(item.content))
        if key in parsed_by_key:
            duplicate_excel_rows += 1
        parsed_by_key[key] = item

    inserted = 0
    updated = 0
    unchanged = 0
    duplicate_db_rows = 0
    type_counter = Counter()
    category_counter = Counter()

    await init_db()

    async with AsyncSessionLocal() as session:
        result = await session.execute(select(Question))
        existing_by_key: dict[tuple[str, str], list[Question]] = {}
        for question in result.scalars().all():
            key = (question.question_type, normalize_content(question.content))
            existing_by_key.setdefault(key, []).append(question)

        for item in parsed_by_key.values():
            key = (item.question_type, normalize_content(item.content))
            matches = existing_by_key.get(key, [])
            type_counter[item.question_type] += 1
            category_counter[item.category or "未分类"] += 1

            if not matches:
                inserted += 1
                if apply:
                    session.add(
                        Question(
                            question_type=item.question_type,
                            content=item.content,
                            category=item.category,
                            options=item.options,
                            answer=item.answer,
                            explanation=item.explanation or None,
                            difficulty_label=None,
                            difficulty=1,
                            is_active=True,
                        )
                    )
                continue

            if len(matches) > 1:
                duplicate_db_rows += len(matches) - 1

            row_changed = False
            for match in matches:
                desired_explanation = item.explanation or None
                changed = any(
                    [
                        match.content != item.content,
                        match.options != item.options,
                        match.answer != item.answer,
                        (match.explanation or None) != desired_explanation,
                        match.category != item.category,
                        match.is_active is not True,
                    ]
                )
                if changed:
                    row_changed = True
                    if apply:
                        match.content = item.content
                        match.options = item.options
                        match.answer = item.answer
                        match.explanation = desired_explanation
                        match.category = item.category
                        match.is_active = True

            if row_changed:
                updated += 1
            else:
                unchanged += 1

        if apply:
            await session.commit()
        else:
            await session.rollback()

    await engine.dispose()
    return {
        "read_rows": len(parsed_questions),
        "unique_rows": len(parsed_by_key),
        "duplicate_excel_rows": duplicate_excel_rows,
        "inserted": inserted,
        "updated": updated,
        "unchanged": unchanged,
        "duplicate_db_rows": duplicate_db_rows,
        "type_counts": dict(type_counter),
        "category_counts": dict(category_counter),
    }


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("excel_path", nargs="?", default=os.getenv("EXCEL_PATH") or DEFAULT_EXCEL_PATH)
    parser.add_argument("--apply", action="store_true", help="Write changes to the database")
    return parser.parse_args()


async def main() -> None:
    args = parse_args()
    if not os.path.exists(args.excel_path):
        raise FileNotFoundError(args.excel_path)

    stats = await sync_questions(args.excel_path, apply=args.apply)
    mode = "APPLY" if args.apply else "DRY_RUN"
    print(f"[{mode}] {args.excel_path}")
    for key, value in stats.items():
        print(f"  {key}: {value}")


if __name__ == "__main__":
    asyncio.run(main())
