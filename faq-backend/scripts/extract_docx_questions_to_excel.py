"""Extract question/answer/explanation items from a DOCX into the question-bank Excel format."""

from __future__ import annotations

import re
from pathlib import Path

import docx
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill


DOCX_PATH = Path(
    r"C:\Users\86135\Documents\WeChat Files\wxid_stm8la2kv7jr22\FileStorage\File"
    r"\2026-05\题库-33-48toveeko.docx"
)
TEMPLATE_PATH = Path(
    r"C:\Users\86135\Documents\WeChat Files\wxid_stm8la2kv7jr22\FileStorage\File"
    r"\2026-04\题库2.0_toveeko_v3.xlsx"
)
OUTPUT_PATH = Path(r"D:\project\分销系统\分销系统\outputs\题库-33-48提取版.xlsx")

HEADERS = ["题目", "正确答案", "选项A", "选项B", "选项C", "选项D", "选项E", "选项F", "解析", "题型分类"]


def clean_text(value: str) -> str:
    return re.sub(r"\s+", " ", value.strip())


def is_question_marker(value: str) -> bool:
    return bool(re.match(r"^\d+\.\s*题目$", value.strip()))


def is_answer_label(value: str) -> bool:
    return bool(re.fullmatch(r"(?:\d+\.\s*)?答案[:：]", value.strip()))


def is_explanation_label(value: str) -> bool:
    return bool(re.fullmatch(r"@?\s*(?:\d+\.\s*)?解析[:：]", value.strip()))


def split_options(option_text: str) -> dict[str, str]:
    text = clean_text(option_text)
    pattern = re.compile(r"([A-F])\.\s*")
    matches = list(pattern.finditer(text))
    options: dict[str, str] = {}
    for index, match in enumerate(matches):
        label = match.group(1)
        start = match.end()
        end = matches[index + 1].start() if index + 1 < len(matches) else len(text)
        options[label] = text[start:end].strip()
    return options


def parse_docx(path: Path) -> list[dict[str, str]]:
    doc = docx.Document(path)
    lines = [paragraph.text.strip() for paragraph in doc.paragraphs if paragraph.text.strip()]
    items: list[dict[str, str]] = []

    index = 0
    while index < len(lines):
        if not is_question_marker(lines[index]):
            index += 1
            continue

        index += 1
        question_parts = []
        while index < len(lines) and lines[index] != "选项：":
            question_parts.append(lines[index])
            index += 1

        if index >= len(lines):
            break

        index += 1
        option_parts = []
        while index < len(lines):
            current = lines[index]
            next_value = lines[index + 1] if index + 1 < len(lines) else ""
            next_next_value = lines[index + 2] if index + 2 < len(lines) else ""
            typo_answer_label = (
                current == "选项："
                and re.fullmatch(r"[A-F]+", next_value.strip() or "")
                and is_explanation_label(next_next_value)
            )
            if is_answer_label(current) or typo_answer_label:
                if typo_answer_label:
                    index += 1
                break
            option_parts.append(current)
            index += 1

        if index >= len(lines):
            break

        if is_answer_label(lines[index]):
            index += 1

        answer = clean_text(lines[index]) if index < len(lines) else ""
        index += 1

        if index < len(lines) and is_explanation_label(lines[index]):
            index += 1

        explanation_parts = []
        while index < len(lines) and not is_question_marker(lines[index]):
            explanation_parts.append(lines[index])
            index += 1

        question = clean_text(" ".join(question_parts))
        options = split_options(" ".join(option_parts))
        explanation = "\n".join(part.strip() for part in explanation_parts).strip()

        items.append(
            {
                "题目": question,
                "正确答案": answer,
                "选项A": options.get("A", ""),
                "选项B": options.get("B", ""),
                "选项C": options.get("C", ""),
                "选项D": options.get("D", ""),
                "选项E": options.get("E", ""),
                "选项F": options.get("F", ""),
                "解析": explanation,
                "题型分类": "",
            }
        )

    return items


def build_workbook(items: list[dict[str, str]]) -> None:
    workbook = openpyxl.load_workbook(TEMPLATE_PATH)
    worksheet = workbook["题库"] if "题库" in workbook.sheetnames else workbook.active
    worksheet.title = "题库"

    if worksheet.max_row > 1:
        worksheet.delete_rows(2, worksheet.max_row - 1)

    for col_index, header in enumerate(HEADERS, start=1):
        cell = worksheet.cell(row=1, column=col_index, value=header)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row_index, item in enumerate(items, start=2):
        for col_index, header in enumerate(HEADERS, start=1):
            cell = worksheet.cell(row=row_index, column=col_index, value=item.get(header, ""))
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    widths = {
        "A": 58,
        "B": 12,
        "C": 32,
        "D": 32,
        "E": 32,
        "F": 32,
        "G": 24,
        "H": 24,
        "I": 72,
        "J": 14,
    }
    for col_letter, width in widths.items():
        worksheet.column_dimensions[col_letter].width = width

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = f"A1:J{len(items) + 1}"

    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(OUTPUT_PATH)


def validate(items: list[dict[str, str]]) -> list[str]:
    issues = []
    for idx, item in enumerate(items, start=1):
        if not item["题目"]:
            issues.append(f"{idx}: missing question")
        if not re.fullmatch(r"[A-F]+", item["正确答案"] or ""):
            issues.append(f"{idx}: unusual answer {item['正确答案']!r}")
        option_count = sum(1 for key in ["选项A", "选项B", "选项C", "选项D", "选项E", "选项F"] if item[key])
        if option_count < 2:
            issues.append(f"{idx}: only {option_count} options")
        if not item["解析"]:
            issues.append(f"{idx}: missing explanation")
    return issues


def main() -> None:
    items = parse_docx(DOCX_PATH)
    issues = validate(items)
    build_workbook(items)
    print(f"output={OUTPUT_PATH}")
    print(f"rows={len(items)}")
    print(f"issues={len(issues)}")
    for issue in issues[:30]:
        print(f"issue: {issue}")


if __name__ == "__main__":
    main()
